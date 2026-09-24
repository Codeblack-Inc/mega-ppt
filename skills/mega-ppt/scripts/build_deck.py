# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0", "pymupdf"]
# ///
"""deck.json → .pptx. Usage: uv run build_deck.py deck.json -o out.pptx

A content slide = header (kicker · headline · lead) + body grid of panels.
Panels are plain functions in PANELS; full-bleed slides in FULL. Spec: references/layouts.md.
"""
import argparse
import io
import json
import math
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION,
                             XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

THEME = {
    "font": "Pretendard",
    "paper": "#FFFFFF",
    "ink": "#17152B",
    "text": "#2B293B",
    "muted": "#68657B",
    "rule": "#DED9ED",
    "soft": "#F6F5FA",
    "accent": "#B5452D",
    "title_style": "tab",  # panel titles: "tab" | "bar" | "line"
}
W, H = 13.333, 7.5
M = 0.5          # side margin
GAP = 0.16       # grid gutter
BODY_SIZE = 11
BODY_MIN = 9     # floor panels pass for body text
MIN_SIZE = 8     # hard floor
TITLE_H = 0.40   # panel title tab + gap above the body
CAP_H = 0.30     # figure caption strip
HEAD_BEFORE = 4  # pt above a □ heading paragraph
ROMAN = "Ⅰ Ⅱ Ⅲ Ⅳ Ⅴ Ⅵ Ⅶ Ⅷ Ⅸ Ⅹ".split()
SERIES = ["ink", "grey1", "grey2", "rule"]  # chart series order; accent only via highlight
# line weights (pt)
STRONG, FRAME, BORDER, HAIR = 2.0, 1.5, 0.75, 0.5

T = dict(THEME)
CTX = {"slide": 0, "base": Path("."), "fig": 0}
WARN = []

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
EMPH = re.compile(r"\*\*(.+?)\*\*")
HEADING = re.compile(r"\*\*[^*]+\*\*")


def body_max(w):
    """Largest size body text may grow to in a box of width w."""
    return 13 if w >= 4 else 12


# ------------------------------------------------------------------ primitives

def col(c):
    return T.get(c, c)


def rgb(c):
    return RGBColor.from_string(col(c).lstrip("#"))


def mix(a, b, t):
    """Blend colour a toward b by t (0 = a, 1 = b)."""
    ha, hb = col(a).lstrip("#"), col(b).lstrip("#")
    return "#%02X%02X%02X" % tuple(round(int(ha[i:i + 2], 16) * (1 - t) + int(hb[i:i + 2], 16) * t)
                                   for i in (0, 2, 4))


def tint(c, a):
    """Mix colour c with white; a=0.9 → 90% white."""
    return mix(c, "#FFFFFF", a)


def derive():
    """Tokens computed from the theme, so a deck that sets only accent/ink still gets a coherent
    palette. A deck theme may override any of them."""
    for k, v in (("ink2", tint("ink", 0.18)), ("grey1", mix("ink", "rule", 0.6)),
                 ("grey2", mix("ink", "rule", 0.82)), ("line", mix("rule", "ink", 0.22)),
                 ("head", mix("soft", "rule", 0.5)), ("accent_soft", tint("accent", 0.9)),
                 ("on_dark", tint("ink", 0.7)), ("on_dark_rule", tint("ink", 0.3)),
                 ("accent_on_dark", tint("accent", 0.55))):
        T.setdefault(k, v)


derive()


def warn(msg):
    WARN.append(f"slide {CTX['slide']}: {msg}")


def plain(t):
    return EMPH.sub(r"\1", str(t))


def _cw(ch):
    """Approx glyph advance in em. ponytail: width table, swap for real font metrics if fit misjudges."""
    o = ord(ch)
    if ch == " ":
        return 0.28
    if 0xAC00 <= o <= 0xD7A3 or 0x1100 <= o <= 0x11FF or 0x3130 <= o <= 0x318F \
            or 0x2E80 <= o <= 0x9FFF or 0xFF00 <= o <= 0xFFEF:
        return 0.94
    if ch.isupper() or ch in "mwMW%@&":
        return 0.66
    if ch.isdigit():
        return 0.57
    if ch in ".,:;!|il'\"()[]-·/":
        return 0.3
    return 0.53


def text_w(t, size, bold=False):
    return sum(_cw(c) for c in plain(t)) * size / 72 * (1.05 if bold else 1)


def n_lines(t, size, width, bold=False):
    maxw = width / (size / 72) / (1.05 if bold else 1)
    if maxw <= 0:  # no room at all: report "never fits" instead of looping
        return 99
    total = 0
    for para in plain(t).split("\n"):
        lines, cur = 1, 0.0
        for word in re.split(r"(\s+)", para):
            wl = sum(_cw(c) for c in word)
            if cur + wl > maxw and cur > 0 and not word.isspace():
                lines += 1
                cur = wl
            else:
                cur += wl
            while cur > maxw:  # word longer than the line
                lines += 1
                cur -= maxw
        total += lines
    return total


# paragraph levels: 0 plain · 1 ○ item · 2 - sub-item · 3 □ heading (item wholly **bold**) · 4 ※ note
LVL_SIZE = {0: 0, 1: 0, 2: -1, 3: 0.5, 4: -1}


def _paras(content, bullets):
    items = [content] if isinstance(content, (str, int, float)) else content
    out = []
    for it in items:
        s = str(it)
        if not bullets:
            out.append((s, 0))
        elif s.startswith("- "):
            out.append((s[2:], 2))
        elif s.startswith("※"):
            out.append((s, 4))
        elif HEADING.fullmatch(s):
            out.append((s[2:-2], 3))
        else:
            out.append((s, 1))
    return out


def _indent(size, lvl):
    return 0 if lvl == 0 else size / 72 * (2.5 if lvl == 2 else 1.25)


def text_height(content, size, w, bullets=False, bold=False, line=1.15, gap=4):
    ps = _paras(content, bullets)
    h = sum(n_lines(t, size + LVL_SIZE[lvl], w - _indent(size, lvl), bold or lvl == 3)
            * (size + LVL_SIZE[lvl]) * line * 1.2 / 72 for t, lvl in ps)
    heads = sum(1 for i, (_, lvl) in enumerate(ps) if i and lvl == 3)
    return h + gap / 72 * (len(ps) - 1) + heads * HEAD_BEFORE / 72


def _check(need, h, size, name, underfill=True, grown=False):
    if need > h * 1.04:
        warn(f"{name!r} overflows its box even at {size}pt — shorten the text")
    elif underfill and (h > 1.2 and need < h * 0.5 if grown else h > 1.5 and need < h * 0.3):
        warn(f"{name!r} fills only {need / h:.0%} of its box — add content or lower the row 'h'")


def fit_size(content, size, w, h, bullets=False, bold=False, min_size=MIN_SIZE, line=1.15,
             gap=4, name="text", underfill=True, max_size=None, quiet=False):
    """Shrink until the text fits. With max_size, first grow toward it while the text still fits
    in 90% of the box (the margin absorbs fallback-font width)."""
    def need(sz):
        return text_height(content, sz, w, bullets, bold, line, gap)
    if max_size:
        while size + 0.5 <= max_size and need(size + 0.5) <= h * 0.9:
            size += 0.5
    while size > min_size and need(size) > h:
        size -= 0.5
    if not quiet:
        _check(need(size), h, size, name, underfill, bool(max_size))
    return size


def share(jobs, size, max_size=None, name="text", underfill=True, bullets=False, bold=False,
          min_size=BODY_MIN, line=1.15, gap=4):
    """One font size for sibling boxes: the largest that fits every (content, w, h) job.
    Warns once if a box overflows, or if even the fullest box stays under half full."""
    jobs = [(c, w, h) for c, w, h in jobs if c]
    if not jobs:
        return size
    s = min(fit_size(c, size, w, h, bullets, bold, min_size, line, gap, max_size=max_size,
                     quiet=True) for c, w, h in jobs)
    fill = [(text_height(c, s, w, bullets, bold, line, gap), h) for c, w, h in jobs]
    worst = max(fill, key=lambda f: f[0] / f[1])
    _check(*worst, s, name, underfill, grown=bool(max_size))
    return s


def _set_font(run, size, color, bold):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = T["font"]
    rpr = run._r.get_or_add_rPr()  # python-pptx sets only latin; Hangul needs <a:ea>
    if size >= 20:  # tighten display type: -2% (-3% from 32pt), in 1/100 pt
        rpr.set("spc", str(int(-(3 if size >= 32 else 2) * size)))
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", T["font"])


BULLET = {1: "○", 2: "-", 3: "□"}


def _bullet(p, lvl, size, color):
    """Korean report hierarchy: □ heading · ○ item · - sub-item · ※ note (hanging, no glyph)."""
    pPr = p._p.get_or_add_pPr()
    mar = _indent(size, lvl)
    pPr.set("marL", str(int(mar * 914400)))
    pPr.set("indent", str(-int((size * 1.1 / 72 if lvl == 2 else mar) * 914400)))
    if lvl == 4:
        etree.SubElement(pPr, qn("a:buNone"))
        return
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr"), val=col(color).lstrip("#"))
    etree.SubElement(pPr, qn("a:buSzPct"), val="90000")
    etree.SubElement(pPr, qn("a:buFont"), typeface=T["font"])
    etree.SubElement(pPr, qn("a:buChar"), char=BULLET[lvl])


def text(s, box, content, size=BODY_SIZE, color="text", bold=False, align="l", anchor="t",
         bullets=False, fit=True, min_size=MIN_SIZE, line=1.15, gap=4, emph="accent",
         name="text", wrap=True, grow=None, spread=False):
    """Text box with **emphasis** markup, the □/○/-/※ hierarchy and fit-to-box.
    grow=max pt lets the text grow to fill the box; spread=True puts spare height before each ○/□
    group (a body with a single group is centred instead)."""
    x, y, w, h = box
    if fit:
        size = fit_size(content, size, w, h, bullets, bold, min_size, line, gap, name,
                        underfill=anchor == "t" or bool(grow), max_size=grow)
    ps = _paras(content, bullets)
    extra = 0
    starts = [i for i, (_, lvl) in enumerate(ps) if i and lvl not in (2, 4)]  # ○/□ group starts
    if spread and anchor == "t":
        need = text_height(content, size, w, bullets, bold, line, gap)
        if need < h * 0.85:  # spare height goes before each group, so '- ' children stay tight
            if starts:
                extra = min(14, (h * 0.9 - need) * 72 / len(starts))
            else:
                anchor = "m"
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    glyph = "ink" if color == "text" else color
    for i, (t, lvl) in enumerate(ps):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = line
        before = (HEAD_BEFORE if lvl == 3 and i else 0) + (extra if i in starts else 0)
        if before:
            p.space_before = Pt(before)
        p.space_after = Pt(gap)
        sz = size + LVL_SIZE[lvl]
        if lvl:
            _bullet(p, lvl, size, glyph if lvl != 2 else color)
        c = "muted" if lvl == 4 else "ink" if lvl == 3 and color == "text" else color
        for j, seg in enumerate(t.split("\n")):
            if j:
                p.add_line_break()
            for k, part in enumerate(EMPH.split(seg)):
                if part:
                    r = p.add_run()
                    r.text = part
                    _set_font(r, sz, emph if k % 2 else c, bold or lvl == 3 or bool(k % 2))
    return tb


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh._element.remove(sh._element.find(qn("p:style")))  # no theme shadow/effects
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def line(s, x1, y1, x2, y2, color="rule", lw=0.75, arrow=False, dash=False, arrow_start=False):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c._element.remove(c._element.find(qn("p:style")))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    if arrow_start:  # schema order: headEnd before tailEnd
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if arrow:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def hline(s, x, y, w, color="rule", lw=0.75):
    return line(s, x, y, x + w, y, color, lw)


def box(s, x, y, w, h, hl=False, fill=None):
    """A grid cell: 0.75pt border and a 1.5pt ink top edge (accent when highlighted)."""
    rect(s, x, y, w, h, fill, "accent" if hl else "line", lw=1.25 if hl else BORDER)
    hline(s, x, y, w, "accent" if hl else "ink", FRAME)


def cell_head(s, x, y, w, h, title, hl=False, num=None, icon_name=None, size=12):
    """Header strip of a cell: head fill (accent when highlighted), optional ink number cell or
    icon, bold title."""
    rect(s, x, y, w, h, "accent" if hl else "head")
    tx = x + 0.12
    if num is not None:
        nc = rect(s, x, y, 0.38, h, "ink")
        shape_text(nc, num, 11, "paper")
        tx = x + 0.5
    elif icon_name:
        icon(s, icon_name, x + 0.12, y + h / 2 - 0.13, 0.26, "paper" if hl else "ink")
        tx = x + 0.5
    text(s, (tx, y, x + w - tx - 0.1, h), title, size, "paper" if hl else "ink", bold=True,
         anchor="m", min_size=9, line=1.05, name="cell title")


def shape_text(sh, t, size, color, bold=True, align="c", box=None):
    """Text inside a shape. box=(w, h) usable area → shrink to fit (never wraps mid-word)."""
    if box:
        words = max((wd for ln in plain(t).split("\n") for wd in ln.split()), key=len, default="")
        while size > 7 and (text_w(words, size, bold) > box[0] * 0.85 or  # margin for fallback fonts
                            text_height(t, size, box[0], bold=bold, line=1.0) > box[1]):
            size -= 0.5
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = ALIGN[align]
    for j, seg in enumerate(str(t).split("\n")):
        if j:
            p.add_line_break()
        r = p.add_run()
        r.text = plain(seg)
        _set_font(r, size, color, bold)


# ------------------------------------------------------------------ panels
# Each panel: f(slide, (x, y, w, h), spec). Optional spec["title"] is drawn by panel().

def _body(s, box, content, size, bullets, name, gap=4):
    """A sibling cell body at a size already shared across the row (see share())."""
    return text(s, box, content, size, bullets=bullets, gap=gap, fit=False, name=name, spread=True)

def p_bullets(s, b, p):
    x, y, w, h = b
    items, bl, gap = p["items"], p.get("style") != "plain", p.get("gap", 5)
    parts = [items]
    if p.get("cols") == 2 and len(items) > 1:  # split at the □ heading nearest the middle
        cut = [i for i, t in enumerate(items) if i and HEADING.fullmatch(str(t))] \
            or [i for i, t in enumerate(items) if i and not str(t).startswith(("- ", "※"))] or [1]
        k = min(cut, key=lambda i: abs(i - len(items) / 2))
        parts = [items[:k], items[k:]]
    cw = (w - 0.3 * (len(parts) - 1)) / len(parts)
    name = p.get("title", "bullets")
    size = share([(c, cw, h) for c in parts], p.get("size", BODY_SIZE), body_max(cw), name,
                 bullets=bl, gap=gap)
    for i, c in enumerate(parts):
        if i:
            line(s, x + i * (cw + 0.3) - 0.15, y, x + i * (cw + 0.3) - 0.15, y + h, "rule", HAIR)
        _body(s, (x + i * (cw + 0.3), y, cw, h), c, size, bl, name, gap)


def p_callout(s, b, p):
    """Default: the 시사점 box (ink label cell + bold text). 'dark': ink slab, once per deck.
    'soft': quiet side note."""
    x, y, w, h = b
    st, t = p.get("style"), p["text"]
    size, top, color, emph, tx = p.get("size", 12.5), 14, "ink", "accent", x + 0.2
    if st == "soft":
        rect(s, x, y, w, h, "soft")
        size, top, color = p.get("size", 11), 12, "text"
    elif st == "dark":
        rect(s, x, y, w, h, "ink")
        t, tx, color, emph = "**⇒** " + t, x + 0.25, "paper", "accent_on_dark"
    else:
        rect(s, x, y, w, h, "paper", "ink", lw=1.25)
        lab = p.get("label", "시사점")
        if lab:
            shape_text(rect(s, x, y, 1.1, h, "ink"), lab, 10.5, "paper", box=(1.0, h))
            tx = x + 1.28
    bw, bh, bold = max(0.5, x + w - 0.2 - tx), h - 0.12, st != "soft"
    # a centred slab: grow to fill, but never an underfill warning
    sz = fit_size(t, size, bw, bh, bold=bold, min_size=9 if st == "soft" else 10, max_size=top,
                  underfill=False, name="callout")
    text(s, (tx, y + 0.06, bw, bh), t, sz, color, bold=bold, anchor="m",
         align=p.get("align", "l"), emph=emph, fit=False)


NUMERIC = re.compile(r"[+\-−±~]?[\d.,]+(\s*(%p?|배|×|x|[BMK]|억|만|천|원|건|곳|명|개|점|편|회))*|-|")
HEAT = {"상": ("ink2", "paper"), "중": ("grey2", "ink"), "하": ("head", "ink")}  # risk levels


def p_table(s, b, p):
    """Report table: head-filled header, hairline grid, right-aligned numeric columns, optional
    rowhead column, highlight row / highlight_col, total row, heat (상/중/하) columns."""
    x, y, w, h = b
    if p.get("unit"):  # titled tables get it in the panel note instead
        text(s, (x, y, w, 0.22), f"(단위: {p['unit']})", 8.5, "muted", align="r", fit=False)
        y, h = y + 0.22, h - 0.22
    cols, rows = p.get("columns"), p["rows"]
    data = ([cols] if cols else []) + rows
    nr, nc = len(data), max(len(r) for r in data)
    hdr = 1 if cols else 0
    val = [[plain(r[c]).strip() if c < len(r) else "" for c in range(nc)] for r in data]
    rel = p.get("widths") or [1] * nc
    cws = [w * r / sum(rel) for r in rel]
    rowhead, total = p.get("rowhead", False), p.get("total", False)
    hl, hc, heat = p.get("highlight"), p.get("highlight_col"), set(p.get("heat", []))
    dark = p.get("header") == "dark"
    body = val[hdr:nr - 1] if total else val[hdr:]
    num = [not (rowhead and c == 0) and any(r[c] not in ("", "-") for r in body)
           and all(NUMERIC.fullmatch(r[c]) for r in body) for c in range(nc)]
    left = [(rowhead and c == 0) or any(len(r[c]) > 14 for r in body) for c in range(nc)]

    def heights(sz):
        out = []
        for r in range(nr):
            need = max((sz + 6) * 1.2 / 72 if val[r][c] in MARKS and r >= hdr else
                       n_lines(val[r][c], sz, cws[c] - 0.2, r < hdr or (rowhead and c == 0))
                       * sz * 1.25 / 72 for c in range(nc))
            out.append(max(need + 0.1, 0.36 if r < hdr else 0.30))
        return out

    def words_fit(sz):  # growing never breaks a word in its cell (10% margin for fallback fonts)
        return all(text_w(wd, sz, r < hdr or (rowhead and c == 0)) <= (cws[c] - 0.2) * 0.9
                   for r in range(nr) for c in range(nc) for wd in val[r][c].split())
    size, top = p.get("size", 10.5), max(12, p.get("size", 10.5))
    while size + 0.5 <= top and sum(heights(size + 0.5)) <= h * 0.9 and words_fit(size + 0.5):
        size += 0.5
    while size > MIN_SIZE and sum(heights(size)) > h:
        size -= 0.5
    rh = heights(size)
    if sum(rh) > h * 1.04:
        warn(f"table {p.get('title', '')!r} too tall ({sum(rh):.1f}in > {h:.1f}in) — cut rows")
    if p.get("stretch", True) and sum(rh[hdr:]):  # body rows stretch, ≤0.65in a row on average
        k = (min(h, 0.65 * nr) - sum(rh[:hdr])) / sum(rh[hdr:])
        rh = rh[:hdr] + [r * max(k, 1) for r in rh[hdr:]]
    if h > 2.0 and sum(rh) < h * 0.8:
        warn(f"table {p.get('title', '')!r} covers only {sum(rh) / h:.0%} of its panel — add "
             "rows/columns or a callout row, or lower the row h")
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(rh)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
    for i, cw in enumerate(cws):
        tbl.columns[i].width = Inches(cw)

    tab = bool(p.get("title"))  # panel() lifted us onto the title rule: it is our top edge

    def edge(r):  # horizontal rule above row r (r == nr: table bottom)
        if r == 0 and tab:
            return None
        if r in (0, nr):
            return "ink", FRAME
        if r == hdr or (total and r == nr - 1):
            return "ink", BORDER
        return "line", HAIR
    for r in range(nr):
        tbl.rows[r].height = Inches(rh[r])
        head = r < hdr
        is_hl = hl is not None and r - hdr == hl
        is_total = total and r == nr - 1
        for c in range(nc):
            v, raw = val[r][c], str(data[r][c]) if c < len(data[r]) else ""
            cell = tbl.cell(r, c)
            rhc = rowhead and c == 0 and not head
            hot = HEAT.get(v) if c in heat and not head else None
            in_hc = hc == c
            fill = (("accent" if in_hc else "ink" if dark else "head") if head else
                    hot[0] if hot else "head" if is_total else
                    "accent_soft" if is_hl or in_hc else "soft" if rhc else "paper")
            gap = ("paper", FRAME)  # heat tiles: a paper gutter between equal neighbours
            _cell_borders(cell, L=gap if hot else ("rule", HAIR) if c else None,
                          R=gap if hot else ("rule", HAIR) if c < nc - 1 else None,
                          T=gap if hot and r > hdr else edge(r),
                          B=gap if hot and r < nr - 1 else edge(r + 1))
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.12 if num[c] else 0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para.alignment = (PP_ALIGN.RIGHT if num[c] else
                              PP_ALIGN.LEFT if left[c] and not (head and c) else PP_ALIGN.CENTER)
            if v in MARKS and not head:
                para.alignment = PP_ALIGN.CENTER
                run = para.add_run()
                run.text = v
                _set_font(run, size + 6, "accent" if "**" in raw else MARKS[v], False)
                continue
            color = ("paper" if head and (dark or in_hc) else hot[1] if hot else
                     "ink" if head or rhc or is_hl or is_total else "text")
            bold = head or rhc or is_hl or is_total or bool(hot)
            for k, part in enumerate(EMPH.split(raw)):
                if part:
                    run = para.add_run()
                    run.text = part
                    _set_font(run, size, "accent" if k % 2 and not head else color,
                              bold or bool(k % 2))
    for r1, c1, r2, c2 in p.get("merge", []):  # data coords, header row = 0 when columns given
        tbl.cell(r1, c1).merge(tbl.cell(r2, c2))
        if r2 > r1 and not (rowhead and c1 == 0):  # merged text runs from the top
            tbl.cell(r1, c1).vertical_anchor = MSO_ANCHOR.TOP


MARKS = {"●": "ink", "✓": "ink", "O": "ink", "◐": "grey1", "○": "grey1", "✗": "grey1",
         "△": "grey1", "X": "grey1"}


def _cell_borders(cell, L=None, R=None, T=None, B=None):
    """Per-side cell borders; each side is None (no line) or (colour, width pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for i, (tag, side) in enumerate((("a:lnL", L), ("a:lnR", R), ("a:lnT", T), ("a:lnB", B))):
        ln = etree.Element(qn(tag), w=str(int((side[1] if side else 0.5) * 12700)))
        if side:
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=col(side[0]).lstrip("#"))
        else:
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)  # schema order: lnL lnR lnT lnB, then fill


CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.DOUGHNUT,
}
PIE = ["ink", "grey1", "grey2", "rule", "head"]  # doughnut slices after the highlighted one


def _chart_new(s, b, kind, cats, series):
    """A chart with the report defaults: 9pt text, no title or legend, 1pt ink category axis."""
    x, y, w, h = b
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    ch = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    ch.font.size, ch.font.name = Pt(9), T["font"]
    ch.font.color.rgb = rgb("text")
    etree.SubElement(ch.font._rPr, qn("a:ea"), typeface=T["font"])  # Hangul labels
    ch.has_title = ch.has_legend = False
    if kind != XL_CHART_TYPE.DOUGHNUT:
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb("ink")
        ca.format.line.width = Pt(1)
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.tick_labels.font.size = Pt(9)
    return ch


def _dlabel(dl, size=9.5, color="text", pos=None, sep=None, cat=False, ser=False, nf=None):
    """Style a data label: a series' DataLabels, or one point's DataLabel (value shown)."""
    dl.font.size, dl.font.bold = Pt(size), True
    dl.font.color.rgb = rgb(color)
    if pos is not None:
        dl.position = pos
    if hasattr(dl, "_idx"):  # point label: python-pptx has no show_* flags here
        el = dl._get_or_add_dLbl()
        for tag, on in (("c:showVal", True), ("c:showCatName", cat), ("c:showSerName", ser)):
            el.find(qn(tag)).set("val", "1" if on else "0")
        if nf:  # CT_DLbl order: idx, layout, tx, numFmt, spPr, txPr, …
            nxt = next(e for e in el if e.tag in (qn("c:spPr"), qn("c:txPr"), qn("c:dLblPos"),
                                                  qn("c:showLegendKey")))
            nxt.addprevious(etree.Element(qn("c:numFmt"), formatCode=nf, sourceLinked="0"))
    else:
        el = dl._element
        dl.show_value, dl.show_category_name, dl.show_series_name = True, cat, ser
        if nf:
            dl.number_format, dl.number_format_is_linked = nf, False
    if sep is not None:  # separator follows the show* flags
        etree.SubElement(el, qn("c:separator")).text = sep


def _nice_min(vals):
    """Axis floor for a line chart: below the data by ~60% of its span, on a round number."""
    lo, hi = min(vals), max(vals)
    if hi == lo:  # flat series: float it above the axis
        return lo - max(1, abs(lo) * 0.2)
    step = 10 ** math.floor(math.log10(hi - lo))
    f = math.floor((lo - (hi - lo) * 0.6) / step) * step
    return f if lo < 0 else max(0, f)




def _key(s, x, y, w, names, colors):
    """Series key in series order, right-aligned (LibreOffice lists hbar legends backwards)."""
    cx = x + w
    for nm, c in reversed(list(zip(names, colors))):
        tw = text_w(nm, 9) + 0.06
        cx -= tw
        text(s, (cx, y, tw, 0.2), nm, 9, "text", anchor="m", fit=False)
        rect(s, cx - 0.16, y + 0.045, 0.11, 0.11, c)
        cx -= 0.34


def p_chart(s, b, p):
    x, y, w, h = b
    if p.get("unit"):  # titled charts get it in the panel note instead
        text(s, (x, y, w, 0.22), f"(단위: {p['unit']})", 8.5, "muted", align="r", fit=False)
        y, h = y + 0.22, h - 0.22
    kind = p.get("kind", "bar")
    if kind == "waterfall":
        return _waterfall(s, (x, y, w, h), p)
    if kind == "combo":
        return _combo(s, (x, y, w, h), p)
    series, cats = p["series"], p["categories"]
    n, hl, nf = len(series), p.get("highlight"), p.get("number_format")
    point = n == 1 and kind not in ("line", "pie")  # single-series bars: highlight = category
    lim = len(cats) if point or kind == "pie" else n
    if hl is not None and not (isinstance(hl, int) and 0 <= hl < lim):
        warn(f"chart {p.get('title', '')!r}: highlight {hl!r} out of range 0..{lim - 1}")
        p, hl = {**p, "highlight": None}, None
    greys = iter(SERIES[1:] + ["head"] * n)
    colors = (["grey1" if point else "accent"] if n == 1 and hl is not None else
              ["ink2"] if point else [SERIES[i % 4] for i in range(n)] if hl is None else
              ["accent" if i == hl else next(greys) for i in range(n)])
    if kind == "hbar" and n > 1:  # own key: the renderers disagree on hbar legend order
        _key(s, x, y, w, [sr["name"] for sr in series], colors)
        y, h = y + 0.26, h - 0.26
    ch = _chart_new(s, (x, y, w, h), CHART_TYPES[kind], cats,
                    [(sr["name"], sr["values"]) for sr in series])
    plot = ch.plots[0]
    if kind == "pie":
        return _doughnut(s, (x, y, w, h), ch, p)
    multiline = kind == "line" and n > 1
    labels = p.get("labels", True) and not multiline
    va = ch.value_axis
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.font.size = Pt(9)
    va.has_major_gridlines = not labels
    if labels:
        va.visible = False
    else:
        va.major_gridlines.format.line.color.rgb = rgb("rule")
        va.major_gridlines.format.line.width = Pt(HAIR)
    lo = p.get("min", _nice_min([v for sr in series for v in sr["values"] if v is not None])
               if kind == "line" else None)
    if lo is not None:
        va.minimum_scale = lo
    if p.get("max") is not None:
        va.maximum_scale = p["max"]
    if kind == "hbar":
        ch.category_axis.reverse_order = True  # first category on top, like the source table
    if kind == "line":
        va.crosses = XL_AXIS_CROSSES.MINIMUM  # category labels stay at the foot below negatives
        last = len(cats) - 1
        named = 1 < n <= 3 and max(len(sr["name"]) for sr in series) <= 6  # else a legend
        for i, ser in enumerate(plot.series):
            c = colors[i]
            ser.smooth = False
            ser.format.line.color.rgb = rgb(c)
            ser.format.line.width = Pt(2.25 if n == 1 or hl == i else 1.75)
            ser.marker.style = XL_MARKER_STYLE.NONE
            m = ser.points[last].marker
            m.style, m.size = XL_MARKER_STYLE.CIRCLE, 7
            m.format.fill.solid()
            m.format.fill.fore_color.rgb = rgb(c)
            m.format.line.color.rgb = rgb(c)
            for j in (range(last + 1) if labels else [last]):
                end = named and j == last
                _dlabel(ser.points[j].data_label, 9.5, c, XL_LABEL_POSITION.RIGHT if j == last
                        and n > 1 else XL_LABEL_POSITION.ABOVE, "\n" if end else None, ser=end,
                        nf=nf)
        if n > 1 and not named:
            _legend(ch)
        return
    plot.gap_width = 60
    plot.overlap = 100 if kind == "stacked" else -10 if n > 1 else 0
    for i, ser in enumerate(plot.series):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(colors[i])
    if point and hl is not None:
        pt = plot.series[0].points[hl]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb("accent")
    if labels:
        stacked = kind == "stacked"  # labels sit inside the segments
        for i, ser in enumerate(plot.series):
            fg = "paper" if colors[i] in ("ink", "grey1", "accent") else "ink"
            _dlabel(ser.data_labels, 9.5, fg if stacked else "text", XL_LABEL_POSITION.CENTER
                    if stacked else XL_LABEL_POSITION.OUTSIDE_END, nf=nf)
    if n > 1 and kind != "hbar":
        _legend(ch)


def _legend(ch):
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
    ch.legend.font.size = Pt(9)


def _doughnut(s, b, ch, p):
    """Doughnut: 60% hole, highlight slice in accent, category + share on the slice."""
    x, y, w, h = b
    hl, nf = p.get("highlight"), p.get("number_format", "0%")
    plot = ch.plots[0]
    hole = plot._element.find(qn("c:holeSize"))
    if hole is None:
        hole = etree.SubElement(plot._element, qn("c:holeSize"))
    hole.set("val", "60")
    others = iter(PIE * 3)
    for i, pt in enumerate(plot.series[0].points):
        c = "accent" if i == hl else next(others)
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.color.rgb = rgb("paper")
        pt.format.line.width = Pt(1.5)
        _dlabel(pt.data_label, 10, "paper" if c in ("accent", "ink", "grey1") else "ink",
                sep="\n", cat=True, nf=nf)
    if p.get("center"):
        text(s, (x, y + h / 2 - 0.45, w, 0.9), p["center"], 16, "ink", bold=True, align="c",
             anchor="m", line=1.0, fit=False)


def _combo(s, b, p):
    """Bars on the primary axis + series with "type": "line" on a hidden secondary axis. The axes
    are banded (bars below, lines in a band above) so bar and line labels never collide."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    bars = [sr for sr in p["series"] if sr.get("type", "bar") != "line"]
    lines = [sr for sr in p["series"] if sr.get("type") == "line"]
    if not lines or not bars:
        return p_chart(s, b, {**p, "kind": "line" if lines else "bar", "unit": None})
    ch = _chart_new(s, b, XL_CHART_TYPE.COLUMN_CLUSTERED, p["categories"],
                    [(sr["name"], sr["values"]) for sr in bars + lines])
    plot_area = ch._chartSpace.chart.plotArea
    bar = plot_area.find(qn("c:barChart"))
    lc = parse_xml(f'<c:lineChart {nsdecls("c")}><c:grouping val="standard"/>'
                   '<c:varyColors val="0"/></c:lineChart>')
    for ser in bar.findall(qn("c:ser"))[len(bars):]:
        bar.remove(ser)
        inv = ser.find(qn("c:invertIfNegative"))
        if inv is not None:
            ser.remove(inv)
        etree.SubElement(ser, qn("c:smooth"), val="0")
        lc.append(ser)
    for tag, v in (("c:marker", "1"), ("c:axId", "50001"), ("c:axId", "50000")):  # cat, val
        etree.SubElement(lc, qn(tag), val=v)
    bar.addnext(lc)
    axes = [e for e in plot_area if e.tag in (qn("c:valAx"), qn("c:catAx"))]
    axes[-1].addnext(parse_xml(
        f'<c:catAx {nsdecls("c")}><c:axId val="50001"/><c:scaling><c:orientation val="minMax"/>'
        '</c:scaling><c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/>'
        '<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/><c:crossAx val="50000"/>'
        '<c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/>'
        '<c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'))
    axes[-1].addnext(parse_xml(
        f'<c:valAx {nsdecls("c")}><c:axId val="50000"/><c:scaling><c:orientation val="minMax"/>'
        '</c:scaling><c:delete val="0"/><c:axPos val="r"/><c:numFmt formatCode="General" '
        'sourceLinked="1"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
        '<c:tickLblPos val="none"/><c:spPr><a:ln><a:noFill/></a:ln></c:spPr>'
        '<c:crossAx val="50001"/><c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>'
        .replace("<c:valAx ", f'<c:valAx {nsdecls("a")} ')))
    top = max(v or 0 for sr in bars for v in sr["values"])
    lv = [v for sr in lines for v in sr["values"] if v is not None]
    lo, hi = min(lv), max(lv)
    span = (hi - lo) or abs(hi) or 1
    band = {"0": (0, top / 0.52),  # bars: lower 52%; line: 66–88% of the height
            "1": (lo - 0.66 * span / 0.22, lo + 0.34 * span / 0.22)}
    for va in plot_area.findall(qn("c:valAx")):  # by id: python-pptx may pick the secondary
        secondary = va.find(qn("c:axId")).get("val") == "50000"
        va.find(qn("c:delete")).set("val", "0" if secondary else "1")
        grid = va.find(qn("c:majorGridlines"))
        if grid is not None:
            va.remove(grid)
        mn, mx = band["1" if secondary else "0"]
        sc = va.find(qn("c:scaling"))
        etree.SubElement(sc, qn("c:max"), val=f"{mx:.4g}")
        etree.SubElement(sc, qn("c:min"), val=f"{mn:.4g}")
    bar_plot, line_plot = ch.plots[0], ch.plots[1]
    bar_plot.gap_width = 60
    bar_plot.overlap = -10 if len(bars) > 1 else 0
    hl = p.get("highlight")
    if isinstance(hl, int) and not 0 <= hl < len(p["categories"]):
        warn(f"chart {p.get('title', '')!r}: highlight {hl} out of range")
        hl = None
    for i, (ser, spec) in enumerate(zip(bar_plot.series, bars)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(SERIES[1 + i % 3])
        if isinstance(hl, int) and i == 0:
            pt = ser.points[hl]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb("accent")
    lcol = "accent" if hl == "line" else "ink"
    last = len(p["categories"]) - 1
    for ser in line_plot.series:
        ser.format.line.color.rgb = rgb(lcol)
        ser.format.line.width = Pt(2.25)
        ser.marker.style = XL_MARKER_STYLE.NONE
        m = ser.points[last].marker
        m.style, m.size = XL_MARKER_STYLE.CIRCLE, 7
        m.format.fill.solid()
        m.format.fill.fore_color.rgb = rgb(lcol)
        m.format.line.color.rgb = rgb(lcol)
    for plot, specs, color, pos in ((bar_plot, bars, "text", XL_LABEL_POSITION.OUTSIDE_END),
                                    (line_plot, lines, lcol, XL_LABEL_POSITION.ABOVE)):
        for ser, spec in zip(plot.series, specs):
            _dlabel(ser.data_labels, 9.5, color, pos, nf=spec.get("number_format"))
    if len(bars) + len(lines) > 2:
        _legend(ch)


def _fmt(v, nf):
    """Python rendering of a simple Excel number format ("0.0", "#,##0", "0.0%", '0"억"').
    ponytail: one section, quoted literals only as a suffix."""
    core = nf.split(";")[0]
    lit = "".join(re.findall(r'"([^"]*)"', core))
    core = re.sub(r'"[^"]*"', "", core)
    pct = core.rstrip().endswith("%")
    d = len(re.sub(r"[^0#]", "", core.split(".")[1])) if "." in core else 0
    v = v * 100 if pct else v
    return (f"{v:,.{d}f}" if "," in core else f"{v:.{d}f}") + ("%" if pct else "") + lit




def _waterfall(s, b, p):
    """Bridge chart as stacked columns with an invisible base; p["totals"] = absolute bars.
    Labels ride on an invisible sliver stacked on top of each bar, so they sit above it."""
    vals, totals = p["series"][0]["values"], set(p.get("totals", []))
    base, up, down, run = [], [], [], 0
    for i, v in enumerate(vals):
        if i in totals:
            base.append(0), up.append(v), down.append(0)
            run = v
        elif v >= 0:
            base.append(run), up.append(v), down.append(0)
            run += v
        else:
            run += v
            base.append(run), up.append(0), down.append(-v)
    pad = max(abs(v) for v in base + up) * 0.001
    ch = _chart_new(s, b, XL_CHART_TYPE.COLUMN_STACKED, p["categories"],
                    [("base", base), ("증가", up), ("감소", down), ("label", [pad] * len(vals))])
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, 100
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    nf, hl = p.get("number_format", "#,##0"), p.get("highlight")
    b_ser, u_ser, d_ser, l_ser = plot.series
    b_ser.format.fill.background()
    l_ser.format.fill.background()
    for ser, c in ((u_ser, "ink2"), (d_ser, "grey2")):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(c)
    for i, v in enumerate(vals):
        c = "accent" if i == hl else "ink" if i in totals else None
        if c:
            pt = (d_ser if v < 0 and i not in totals else u_ser).points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(c)
        dl = l_ser.points[i].data_label
        dl.text_frame.text = ("−" if v < 0 and i not in totals else "") + _fmt(abs(v), nf)
        dl.position = XL_LABEL_POSITION.INSIDE_BASE
        f = dl.text_frame.paragraphs[0].runs[0].font
        f.size, f.bold = Pt(9.5), True
        f.color.rgb = rgb("accent" if i == hl else "ink")


KPI_NUM = re.compile(r"^([+\-−]?[\d.,]+)(.*)$")


def _kpi_split(v):
    """'99.2%' → ('99.2', '%'); the unit is drawn at half size. Only short, digit-free units."""
    m = KPI_NUM.match(str(v))
    return (m[1], m[2]) if m and len(m[2]) <= 3 and not re.search(r"\d", m[2]) else (str(v), "")


def _kpi_value(s, box, num, unit, size, color):
    tb = s.shapes.add_textbox(*(Inches(v) for v in box))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    for t, sz in ((num, size), (unit, size * 0.5)):
        if t:
            r = para.add_run()
            r.text = t
            _set_font(r, sz, color, True)


def p_kpi(s, b, p):
    """One joined strip: head label row, big number (unit at half size), optional sub line.
    direction 'column': rows with the label cell on the left."""
    x, y, w, h = b
    items = p["items"]
    n = len(items)
    column = p.get("direction") == "column"
    sub = 0.28 if any(it.get("sub") for it in items) else 0
    if column:
        rh, lw = h / n, w * 0.4
        cells = [(x, y + i * rh, w, rh) for i in range(n)]
        vw, ah, cap = w - lw, rh - 0.12, 32
    else:
        cw = w / n
        lh = max(0.34, max(text_height(it["label"], 10, cw - 0.2, bold=True) for it in items)
                 + 0.12)
        cells = [(x + i * cw, y, cw, h) for i in range(n)]
        vw, ah, cap = cw, h - lh - 0.1, 40
    parts = [_kpi_split(it["value"]) for it in items]
    vs = next((v for v in range(cap, 17, -1)
               if v * 1.25 / 72 <= (ah - sub) * 0.8 and all(
                   text_w(nm, v, True) + text_w(u, v * 0.5, True) <= (vw - 0.3) * 0.88
                   for nm, u in parts)), None)
    if vs is None:
        warn("kpi value too wide or strip too low even at 18pt — shorten the value or raise h")
        vs = 18
    vh = vs * 1.25 / 72
    for i, (it, (cx, cy, cw_, ch), (nm, u)) in enumerate(zip(items, cells, parts)):
        hl = p.get("highlight") == i
        if column:
            rect(s, cx, cy, lw, ch, "accent" if hl else "head")
            text(s, (cx + 0.12, cy, max(0.3, lw - 0.2), ch), it["label"], 10.5,
                 "paper" if hl else "ink",
                 bold=True, anchor="m", min_size=9, line=1.1, name="kpi label")
            ax, ay = cx + lw, cy + 0.06
        else:
            rect(s, cx, cy, cw_, lh, "accent" if hl else "head")
            text(s, (cx + 0.1, cy, max(0.3, cw_ - 0.2), lh), it["label"], 10,
                 "paper" if hl else "ink",
                 bold=True, align="c", anchor="m", min_size=8.5, line=1.1, name="kpi label")
            ax, ay = cx, cy + lh + 0.05
        top = ay + (ah - vh - sub) / 2
        _kpi_value(s, (ax + 0.15, top, vw - 0.3, vh), nm, u, vs, "accent" if hl else "ink")
        if it.get("sub"):
            text(s, (ax + 0.1, top + vh, max(0.3, vw - 0.2), sub), it["sub"], 9.5, "muted",
                 align="c",
                 anchor="m", min_size=8, line=1.0, name="kpi sub")
        rect(s, cx, cy, cw_, ch, None, "line", lw=BORDER)
    hline(s, x, y, w, "ink", FRAME)


def p_cards(s, b, p):
    """Grid of cells: head strip (number cell or icon), shared-size body, optional value foot and
    '→ note' strip. A single column turns the head into a left label cell."""
    x, y, w, h = b
    items = p["items"]
    ncol = p.get("cols", len(items))
    nrow = math.ceil(len(items) / ncol)
    g = 0.12
    cw = (w - g * (ncol - 1)) / ncol
    chh = (h - g * (nrow - 1)) / nrow
    numbered = p.get("numbered", True)
    rows = ncol == 1 and len(items) > 1
    pre = 0.5 if numbered or any(it.get("icon") for it in items) else 0.12
    if rows:  # title cell on the left, body to its right
        th = chh
        lw = min(max(1.2, max(text_w(it["title"], 12, True) for it in items) + pre + 0.2), cw * 0.3)
        tjobs = [(it["title"], max(0.3, lw - pre - 0.1), chh - 0.1) for it in items]
    else:
        th = min(0.64, max(0.40, max(text_height(it["title"], 12, cw - 0.62, bold=True)
                                     for it in items) + 0.16))
        lw = cw
        tjobs = [(it["title"], max(0.3, cw - pre - 0.1), th) for it in items]
    ts = share(tjobs, 12, name="card title", underfill=False, bold=True, min_size=10, line=1.05)
    foot = (0.62 if any(it.get("value") for it in items) else 0) + \
           (0.36 if any(it.get("note") for it in items) else 0)
    bx, bw = (lw + 0.14, cw - lw - 0.28) if rows else (0.14, cw - 0.28)
    by, bh = (0.1, chh - 0.2 - foot) if rows else (th + 0.12, chh - th - 0.22 - foot)
    bodies = [it.get("items") or it.get("text") for it in items]
    bl = any("items" in it for it in items)
    bs = share([(bd, bw, bh) for bd in bodies], p.get("size", BODY_SIZE), 12.5,
               p.get("title", "cards"), bullets=bl)
    for i, (it, body) in enumerate(zip(items, bodies)):
        cx, cy = x + (i % ncol) * (cw + g), y + (i // ncol) * (chh + g)
        hl = p.get("highlight") == i
        if hl:
            rect(s, cx, cy, cw, chh, "accent_soft")
        cell_head(s, cx, cy, lw, th, it["title"], hl, i + 1 if numbered and not it.get("icon")
                  else None, it.get("icon"), ts)
        box(s, cx, cy, cw, chh, hl)  # borders over the fills
        if body:
            _body(s, (cx + bx, cy + by, bw, bh), body, bs, "items" in it,
                       f"card {it['title']}")
        fx, fw, fy = (cx + lw, cw - lw, cy + chh - foot) if rows else (cx, cw, cy + chh - foot)
        if it.get("value"):
            hline(s, fx, fy, fw, "line", HAIR)
            vw = text_w(it["value"], 20, True) + 0.1
            if vw > fw * 0.55:
                warn(f"card value {it['value']!r} too long — keep the number, move words to "
                     "value_label")
            text(s, (fx + 0.14, fy, vw, 0.62), it["value"], 20, "accent" if hl else "ink",
                 bold=True, anchor="m", fit=False, wrap=False)
            text(s, (fx + 0.24 + vw, fy, max(0.4, fw - vw - 0.38), 0.62),
                 it.get("value_label", ""), 9.5,
                 "muted", anchor="m", min_size=8, line=1.05, name="card value label")
        if it.get("note"):
            ny = cy + chh - 0.36
            rect(s, fx + 0.02, ny, fw - 0.04, 0.34, "soft")
            text(s, (fx + 0.14, ny, fw - 0.28, 0.34), "→ " + it["note"], 10, "ink", bold=True,
                 anchor="m", min_size=8.5, name="card note")


def p_process(s, b, p):
    """Chevron row over one body box per step; optional label strip (목표 …) at the box foot."""
    x, y, w, h = b
    steps = p["steps"]
    n = len(steps)
    g, ah = 0.04, 0.46
    cw = (w - g * (n - 1)) / n
    bh = h - 0.54
    lab = 0.40 if any(st.get("label") for st in steps) else 0
    ih = bh - 0.24 - lab
    body = [st.get("items") or st.get("text") for st in steps]
    size = share([(bd, cw - 0.38, ih) for bd in body], BODY_SIZE, max_size=12, bullets=True,
                 name="process steps")
    for i, (st, bd) in enumerate(zip(steps, body)):
        cx = x + i * (cw + g)
        hl = p.get("highlight") == i
        sh = rect(s, cx, y, cw, ah, "accent" if hl else "ink2",
                  shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
        sh.adjustments[0] = 0.3
        shape_text(sh, st["title"], 11.5, "paper", box=(cw - 0.4, ah))
        bx, by, bw = cx, y + 0.54, cw - 0.10
        if hl:
            rect(s, bx, by, bw, bh, "accent_soft")
        if bd:
            text(s, (bx + 0.14, by + 0.12, bw - 0.28, ih), bd, size, bullets="items" in st,
                 fit=False, spread=True)
        if st.get("label"):
            rect(s, bx, by + bh - lab, bw, lab, "accent" if hl else "head")
            text(s, (bx + 0.1, by + bh - lab, bw - 0.2, lab),
                 f"{p.get('label_head', '목표')}  {st['label']}", 10, "paper" if hl else "ink",
                 bold=True, align="c", anchor="m", min_size=8, line=1.05, name="step label")
        rect(s, bx, by, bw, bh, None, "accent" if hl else "line", lw=1.25 if hl else BORDER)


def _milestones_of(t):
    """task 'milestone': k | [k, …] | [k, "라벨"] | [[k, "라벨"], …] → [(k, label|None)]."""
    m = t.get("milestone")
    if not m:
        return []
    m = [m] if not isinstance(m, list) or isinstance(m[-1], str) else m
    return [tuple(k) if isinstance(k, list) else (k, None) for k in m]


def p_timeline(s, b, p):
    """Gantt: optional year groups over the period header, owner column, group/task rows,
    milestone diamonds and a dashed today line."""
    x, y, w, h = b
    periods, tasks = p["periods"], p["tasks"]
    lw = p.get("label_width", 1.8)
    ow = p.get("owner_width", 1.0) if any(t.get("owner") for t in tasks) else 0
    gx = x + lw + ow
    gw = (w - lw - ow) / len(periods)
    gh = 0.30 if p.get("groups") else 0
    hh = 0.34
    top = y + gh + hh
    foot = 0.2 if p.get("today") else 0
    rh = (h - gh - hh - foot) / len(tasks)
    bot = top + rh * len(tasks)
    rect(s, x, y, w, gh + hh, "ink")
    k = 0
    for name, span in p.get("groups", []):
        rect(s, gx + k * gw, y, span * gw, gh, "ink2", "ink", lw=1)
        text(s, (gx + k * gw, y, span * gw, gh), name, 9.5, "paper", bold=True, align="c",
             anchor="m", fit=False)
        k += span
    text(s, (x + 0.1, y, lw - 0.2, gh + hh), p.get("label", "구분"), 9.5, "paper", bold=True,
         anchor="m", fit=False)
    if ow:
        text(s, (x + lw, y, ow, gh + hh), p.get("owner_label", "담당"), 9.5, "paper", bold=True,
             align="c", anchor="m", fit=False)
    for i, per in enumerate(periods):
        text(s, (gx + i * gw, y + gh, gw, hh), per, 9.5, "paper", bold=True, align="c",
             anchor="m", fit=False)
    for r, t in enumerate(tasks):
        ry, grp = top + r * rh, t.get("group")
        rect(s, x, ry, w if grp else lw + ow, rh, "head" if grp else "soft")
        text(s, (x + (0.1 if grp else 0.22), ry, lw - (0.15 if grp else 0.27), rh), t["name"], 10,
             "ink" if grp else "text", bold=bool(grp), anchor="m", min_size=8, line=1.0,
             name="task")
        if t.get("owner"):
            text(s, (x + lw + 0.05, ry, ow - 0.1, rh), t["owner"], 9.5, "text", align="c",
                 anchor="m", min_size=8, line=1.0, name="task owner")
    for i in range(len(periods) + 1):
        line(s, gx + i * gw, top, gx + i * gw, bot, "line", HAIR)
    if ow:
        line(s, x + lw, top, x + lw, bot, "line", HAIR)
    for r in range(1, len(tasks)):
        hline(s, x, top + r * rh, w, "line", HAIR)
    hline(s, x, bot, w, "ink", FRAME)
    bh = min(0.22, rh * 0.45)
    for r, t in enumerate(tasks):
        cy = top + r * rh + rh / 2
        spans = t.get("span") or []
        spans = spans if not spans or isinstance(spans[0], list) else [spans]
        th = bh * 0.45 if t.get("group") else bh  # group = thin summary bar over its tasks
        for a, z in spans:  # 1-based inclusive period indices
            bx, bw = gx + (a - 1) * gw + 0.06, (z - a + 1) * gw - 0.12
            rect(s, bx, cy - th / 2, bw, th, "accent" if t.get("highlight") else
                 "ink" if t.get("group") else "ink2")
        if t.get("note") and spans:
            if th >= 0.2 and text_w(t["note"], 8.5) + 0.16 <= bw:
                text(s, (bx, cy - th / 2, bw, th), t["note"], 8.5, "paper", align="c",
                     anchor="m", fit=False)
            else:
                text(s, (bx + bw + 0.06, cy - 0.12, 2.0, 0.24), t["note"], 8.5, "muted",
                     anchor="m", fit=False)
    tx = gx + (p["today"] - 1) * gw if p.get("today") else None
    if tx is not None:  # under the milestone labels
        line(s, tx, top, tx, bot + 0.03, "accent", 1, dash=True)
        text(s, (tx - 0.5, bot + 0.03, 1.0, 0.17), p.get("today_label", "현재"), 8, "accent",
             bold=True, align="c", anchor="m", fit=False)
    for r, t in enumerate(tasks):
        cy = top + r * rh + rh / 2
        ms = _milestones_of(t)
        for k, lab in ms:
            mx = gx + (k - 0.5) * gw
            rect(s, mx - 0.08, cy - 0.08, 0.16, 0.16, "ink", "paper", MSO_SHAPE.DIAMOND, lw=0.75)
            if lab:
                lw_ = text_w(lab, 8.5, True) + 0.1
                end = min([gx + (k2 - 0.5) * gw - 0.1 for k2, _ in ms if k2 > k] + [x + w])
                left = mx + 0.12 + lw_ > end or (  # would hit the next diamond or the today line
                    tx is not None and mx + 0.12 <= tx <= mx + 0.12 + lw_)
                text(s, (mx - 0.12 - lw_ if left else mx + 0.12, cy - 0.12, lw_, 0.24), lab, 8.5,
                     "ink", bold=True, align="r" if left else "l", anchor="m", fit=False)


def p_cycle(s, b, p):
    """Nodes on an ellipse, joined clockwise by arrows, around an optional centre box.
    Nodes are strings or {title, text}."""
    x, y, w, h = b
    nodes = p["nodes"]
    n = len(nodes)
    if n < 3:  # two nodes would draw both arrows on one line through the centre
        sys.exit(f"slide {CTX['slide']}: cycle needs 3–8 nodes; use flow for two")
    rich = any(isinstance(nd, dict) for nd in nodes)
    nw = min(2.1, w * 0.28)
    nh = min(0.95, h * 0.22) if rich else min(0.62, h * 0.16)
    cx, cy = x + w / 2, y + h / 2
    rx, ry = (w - nw) / 2 - 0.05, (h - nh) / 2 - 0.05
    pts = [(cx + rx * math.cos(a), cy + ry * math.sin(a))
           for a in (-math.pi / 2 + 2 * math.pi * i / n for i in range(n))]
    for i in range(n):
        (x1, y1), (x2, y2) = pts[i], pts[(i + 1) % n]
        dx, dy = x2 - x1, y2 - y1
        t0 = min((nw / 2 + 0.08) / abs(dx) if abs(dx) > 1e-6 else 1,
                 (nh / 2 + 0.08) / abs(dy) if abs(dy) > 1e-6 else 1)
        if t0 < 0.5:
            line(s, x1 + t0 * dx, y1 + t0 * dy, x2 - t0 * dx, y2 - t0 * dy, "ink2", 1.5,
                 arrow=True)
    titles = [nd["title"] if isinstance(nd, dict) else nd for nd in nodes]
    size = share([(t, (nw - 0.2) * PAIR_W, nh * (0.45 if rich else 0.9)) for t in titles], 11,
                 max_size=12, bold=True, min_size=9, line=1.05, gap=0, underfill=False,
                 name="cycle node")
    for i, ((nx, ny), nd, t) in enumerate(zip(pts, nodes, titles)):
        hl = p.get("highlight") == i
        rect(s, nx - nw / 2, ny - nh / 2, nw, nh, "accent" if hl else "ink2")
        sub = nd.get("text") if isinstance(nd, dict) else None
        _pair(s, (nx - nw / 2 + 0.1, ny - nh / 2 + 0.04, nw - 0.2, nh - 0.08), t, sub, size,
              "paper", "paper", sub_size=10)  # on_dark is too faint on ink2 at projection
    if p.get("center"):
        ccw = min(2.4, w * 0.3)
        rect(s, cx - ccw / 2, cy - 0.45, ccw, 0.9, "head", "ink", lw=FRAME)
        text(s, (cx - ccw / 2 + 0.1, cy - 0.45, ccw - 0.2, 0.9), p["center"], 12.5, "ink",
             bold=True, align="c", anchor="m", min_size=10, line=1.05, gap=0, name="cycle center")


PAIR_W = 0.9  # share of a _pair box the text is measured against


def _pair(s, box, title, sub, size, fg="ink", sub_fg="muted", align="c", sub_size=None):
    """A bold title over a smaller second line, centred as one block in box (stack items, org
    and diagram nodes)."""
    x, y, w, h = box
    ss = sub_size or max(size - 1.5, 8)
    if not sub:
        text(s, box, title, size, fg, bold=True, align=align, anchor="m", fit=False, line=1.05)
        return
    th = text_height(title, size, w * PAIR_W, bold=True, line=1.05)  # margin: wider fallback font
    sh = text_height(sub, ss, w * PAIR_W, line=1.05)
    cut = y + max(0, (h - th - sh - 0.03) / 2) + th  # title sits on the cut, sub hangs from it
    text(s, (x, y, w, cut - y), title, size, fg, bold=True, align=align, anchor="b", fit=False,
         line=1.05)
    text(s, (x, cut + 0.03, w, y + h - cut - 0.03), sub, ss, sub_fg, align=align, fit=False,
         line=1.05)


def _item(it):
    """stack item / node: 'title\\nsub' string or {title|name, text|role} → (title, sub)."""
    if isinstance(it, dict):
        return it.get("title") or it.get("name"), it.get("text") or it.get("role")
    t, _, sub = str(it).partition("\n")
    return t, sub or None


def p_stack(s, b, p):
    """Layered architecture, top to bottom: a label cell and a row of item boxes per layer.
    One text size is shared by every item of every layer."""
    x, y, w, h = b
    layers = p["layers"]
    g = 0.06
    lh = (h - g * (len(layers) - 1)) / len(layers)
    lw = p.get("label_width", 1.3)
    ih = lh - 0.12
    rows = [[_item(it) for it in ly.get("items", [])] for ly in layers]
    iws = [(w - lw - 0.16 - 0.08 * (len(r) - 1)) / max(len(r), 1) for r in rows]
    cells = [(t, d, iw - 0.16) for r, iw in zip(rows, iws) for t, d in r]
    size = share([([t] + ([d] if d else []), tw * PAIR_W, ih - 0.1) for t, d, tw in cells], 10.5,
                 max_size=12, bold=True, min_size=8, line=1.05, gap=1, underfill=False,
                 name="stack items")
    while size > 8 and any(text_w(wd, size, True) > tw * 0.95  # never break inside a word
                           for t, _, tw in cells for wd in t.split()):
        size -= 0.5
    for i, (ly, r, iw) in enumerate(zip(layers, rows, iws)):
        ly_y = y + i * (lh + g)
        hl = ly.get("highlight")
        rect(s, x, ly_y, lw, lh, "accent" if hl else "ink")
        text(s, (x + 0.08, ly_y, lw - 0.16, lh), ly["label"], 11, "paper", bold=True, align="c",
             anchor="m", min_size=8.5, line=1.05, name="layer label")
        rect(s, x + lw, ly_y, w - lw, lh, "accent_soft" if hl else "soft",
             "accent" if hl else "line", lw=1.25 if hl else BORDER)
        for j, (t, d) in enumerate(r):
            ix = x + lw + 0.08 + j * (iw + 0.08)
            rect(s, ix, ly_y + 0.06, iw, ih, "paper", "accent" if hl else "line",
                 lw=BORDER if hl else HAIR)
            _pair(s, (ix + 0.08, ly_y + 0.06, iw - 0.16, ih), t, d, size)


def p_label(s, b, p):
    x, y, w, h = b
    rect(s, x, y, w, h, "accent" if p.get("highlight") else "ink")
    t = p["text"]
    text(s, (x + 0.04, y + 0.08, w - 0.08, h - 0.16),
         "\n".join(t) if h > w * 2 and " " not in t else t, 12, "paper", bold=True, align="c",
         anchor="m", line=1.25, min_size=9, name="label")
    line(s, x + w, y, x + w, y + h, "ink", FRAME)


def p_arrow(s, b, p):
    x, y, w, h = b
    if p.get("direction") == "down":
        if p.get("text") and w < 1.5:
            warn("arrow text needs a full-width row — put the arrow in its own row or give it 'w'")
        elif p.get("text"):  # a labelled strip says what the arrow means
            rect(s, x, y + (h - 0.34) / 2, w, 0.34, "soft")
            return text(s, (x + 0.2, y + (h - 0.34) / 2, w - 0.4, 0.34), "▼ " + p["text"], 10.5,
                        "ink", bold=True, align="c", anchor="m", min_size=9, name="arrow")
        aw, ah, shp = 0.6, 0.28, MSO_SHAPE.DOWN_ARROW
    else:
        aw, ah, shp = 0.28, 0.6, MSO_SHAPE.RIGHT_ARROW
    return rect(s, x + (w - aw) / 2, y + (h - ah) / 2, aw, ah, "ink2", shape=shp)


def _rotated_label(s, cx, cy, length, t, color="muted", size=9.5):
    tb = text(s, (cx - length / 2, cy - 0.13, length, 0.26), t, size, color, bold=True,
              align="c", anchor="m", fit=False)
    tb.rotation = 270


def _stacked_label(s, x, y, w, h, t):
    """Upright stacked Korean label on an ink2 strip (내/부/역/량), like p_label."""
    shape_text(rect(s, x, y, w, h, "ink2"), "\n".join(t.replace(" ", "")), 10.5, "paper",
               box=(w, h))


def p_matrix(s, b, p):
    """2×2: quadrants (SWOT 등) or positioning map with points."""
    x, y, w, h = b
    ax = 0.35 if p.get("y_label") else 0
    ay = 0.32 if p.get("x_label") else 0
    gx, gy, gw, gh = x + ax, y, w - ax, h - ay
    if p.get("y_label"):
        _rotated_label(s, x + 0.15, gy + gh / 2, gh, p["y_label"])
    if p.get("x_label"):
        text(s, (gx, y + h - 0.26, gw, 0.26), p["x_label"], 9.5, "muted", bold=True, align="c",
             anchor="m", fit=False)
    if p.get("points"):
        return _positioning(s, (gx, gy, gw, gh), p)
    g, ls, lsw = 0.08, 0.28, 0.34
    cl, rl = p.get("col_labels"), p.get("row_labels")
    if cl:  # e.g. 긍정/부정 strips along the top
        gy, gh = gy + ls + g, gh - ls - g
    if rl:  # e.g. 내부/외부 strips down the left
        gx, gw = gx + lsw + g, gw - lsw - g
    qw, qh = (gw - g) / 2, (gh - g) / 2
    for j in range(2):
        if cl:
            lab = rect(s, gx + j * (qw + g), gy - ls - g, qw, ls, "ink2")
            shape_text(lab, cl[j], 9.5, "paper")
        if rl:
            _stacked_label(s, gx - lsw - g, gy + j * (qh + g), lsw, qh, rl[j])
    letters = p.get("letters") or ""
    quads = p["quadrants"][:4]
    bodies = [q.get("items") or q.get("text") for q in quads]
    bl = any("items" in q for q in quads)
    bs = share([(bd, qw - 0.32, qh - 0.64) for bd in bodies], p.get("size", BODY_SIZE), 12.5,
               "matrix", bullets=bl)
    for i, (q, body) in enumerate(zip(quads, bodies)):
        qx, qy = gx + (i % 2) * (qw + g), gy + (i // 2) * (qh + g)
        hl = p.get("highlight") == i
        rect(s, qx, qy, qw, qh, "accent_soft" if hl else None)
        rect(s, qx, qy, qw, 0.42, "head")
        tx = qx + 0.14
        if i < len(letters):
            shape_text(rect(s, qx, qy, 0.42, 0.42, "accent" if hl else "ink"), letters[i], 16,
                       "paper")
            tx = qx + 0.56
        text(s, (tx, qy, qx + qw - tx - 0.1, 0.42), q["title"], 12, "ink", bold=True,
             anchor="m", min_size=10, name="quadrant title")
        if body:
            _body(s, (qx + 0.16, qy + 0.54, qw - 0.32, qh - 0.64), body, bs, "items" in q,
                       q["title"])
        box(s, qx, qy, qw, qh, hl)


def _positioning(s, b, p):
    x, y, w, h = b
    hl = p.get("highlight")
    for i in range(4):
        rect(s, x + (i % 2) * w / 2, y + (i // 2) * h / 2, w / 2, h / 2,
             "accent_soft" if hl == i else None, "line", lw=HAIR)
    line(s, x, y + h, x + w, y + h, "ink", FRAME, arrow=True)
    line(s, x, y + h, x, y, "ink", FRAME, arrow=True)
    for t, bx in (("낮음", (x + 0.02, y + h + 0.03, 0.5, 0.2)),
                  ("높음", (x + w - 0.52, y + h + 0.03, 0.5, 0.2))):
        text(s, bx, t, 8.5, "muted", align="l" if t == "낮음" else "r", fit=False)
    for t, by in (("높음", y), ("낮음", y + h - 0.2)):
        text(s, (x - 0.36, by, 0.3, 0.2), t, 8.5, "muted", align="r", fit=False)
    for i, q in enumerate(p.get("quadrant_labels", [])[:4]):
        qx, qy = x + (i % 2) * w / 2, y + (i // 2) * h / 2
        text(s, (qx + 0.12, qy + 0.08, w / 2 - 0.24, 0.26), q, 9.5,
             "accent" if hl == i else "muted", bold=True, align="r" if i % 2 else "l",
             fit=False)
    for pt in p["points"]:
        on = pt.get("highlight")
        r = 0.13 if on else 0.09
        if abs(pt["x"] - 0.5) < 0.03 or abs(pt["y"] - 0.5) < 0.03:
            warn(f"point {pt['label']!r} sits on the divider — move it off 0.5")
        px, py = (min(0.96, max(0.04, pt[k])) for k in ("x", "y"))  # off the axes and end labels
        cx, cy = x + px * w, y + (1 - py) * h
        rect(s, cx - r, cy - r, 2 * r, 2 * r, "accent" if on else "ink2", shape=MSO_SHAPE.OVAL)
        note = pt.get("note")
        lw_ = max(text_w(pt["label"], 10.5, bool(on)), text_w(note or "", 9)) + r + 0.1
        # label to the right, unless near the right edge or it would cross the divider
        right = px < 0.75 and not (px < 0.5 and cx + lw_ > x + w / 2)
        lx = cx + r + 0.06 if right else cx - r - 0.06 - 2.2
        al = "l" if right else "r"
        text(s, (lx, cy - (0.25 if note else 0.13), 2.2, 0.26), pt["label"], 10.5 if on else 10,
             "accent" if on else "text", bold=bool(on), align=al, anchor="b" if note else "m",
             fit=False)
        if note:  # e.g. '정확도 61% · 비용 1.4×'
            text(s, (lx, cy + 0.01, 2.2, 0.22), note, 9, "muted", align=al, fit=False)


RAMP = ["ink", "ink2", "grey1", "grey2", "head"]  # level fills, top to bottom (paper text on 3)


def p_pyramid(s, b, p, funnel=False):
    """Trapezoid levels forming one pyramid (widens downward) or funnel (narrows), with an
    optional description column: text | items[] per level, plus a tag cell."""
    x, y, w, h = b
    levels = p["levels"]
    n = len(levels)
    body = [lv.get("items") or lv.get("text") for lv in levels]
    tags = [str(lv.get("tag") or "").replace(" · ", "\n") for lv in levels]  # 2 lines at ' · '
    tag = min(1.4, max((text_w(ln, 12, True) / 0.85 + 0.2 for t in tags for ln in t.split("\n")
                        if t), default=0))  # 0.85: margin for wider fallback fonts
    tgs = 12
    while tgs > 9 and any(text_w(ln, tgs, True) > (tag - 0.2) * 0.85
                          for t in tags for ln in t.split("\n")):
        tgs -= 0.5
    sw = w * 0.42 if any(body) else w - tag - 0.5 if tag else w
    g = 0.06
    lh = (h - g * (n - 1)) / n

    def wid(f):  # shape width at depth f (0 = top edge, 1 = bottom edge)
        return sw * (1 - 0.66 * f if funnel else 0.22 + 0.78 * f)
    tops = [y + i * (lh + g) for i in range(n)]
    mids = [wid((ty + lh / 2 - y) / h) * 0.8 for ty in tops]
    vals = any(lv.get("value") for lv in levels)
    ts = share([(lv["title"], mw, lh * (0.4 if vals else 0.8)) for lv, mw in zip(levels, mids)],
               12, max_size=14, bold=True, min_size=9, line=1.05, gap=0, underfill=False,
               name="level title")
    vs = ts + 5
    while vs > 10 and any(text_w(lv.get("value", ""), vs, True) > mw * 0.85
                          for lv, mw in zip(levels, mids)):
        vs -= 0.5
    tx = x + sw + 0.4
    tw = x + w - tx - (tag + 0.14 if tag else 0)
    size = share([(bd, tw, lh - 0.2) for bd in body], 10.5, max_size=body_max(tw), bullets=True,
                 name="level text")
    cx = x + sw / 2
    for i, (lv, bd, by, mw) in enumerate(zip(levels, body, tops, mids)):
        wt, wb = wid((by - y) / h), wid((by + lh - y) / h)
        bw = max(wt, wb)
        hl = p.get("highlight") == i
        k = min(i, len(RAMP) - 1)
        fg = "paper" if hl or k < 3 else "ink"
        sh = rect(s, cx - bw / 2, by, bw, lh, "accent" if hl else RAMP[k],
                  shape=MSO_SHAPE.TRAPEZOID)
        sh.adjustments[0] = abs(wb - wt) / 2 / min(bw, lh)
        if funnel:
            sh.rotation = 180
        if lv.get("value"):  # title sits on the midline, value hangs from it
            text(s, (cx - mw / 2, by, mw, lh / 2 - 0.02), lv["title"], ts, fg, bold=True,
                 align="c", anchor="b", fit=False, line=1.05)
            text(s, (cx - mw / 2, by + lh / 2, mw, lh / 2), lv["value"], vs, fg, bold=True,
                 align="c", fit=False, line=1.0)
        else:
            text(s, (cx - mw / 2, by, mw, lh), lv["title"], ts, fg, bold=True, align="c",
                 anchor="m", fit=False, line=1.05)
        lx, my = cx + wid((by + lh / 2 - y) / h) / 2 + 0.06, by + lh / 2
        if bd or tags[i]:
            line(s, lx, my, tx - 0.12 if bd else x + w - tag - 0.06, my, "ink", HAIR)
        if bd:
            text(s, (tx, by + 0.1, tw, lh - 0.2), bd, size, bullets=isinstance(bd, list),
                 fit=False, spread=True)
        if tags[i]:  # sized to its text, centred on the leader
            tth = min(lh - 0.1, text_height(tags[i], tgs, tag - 0.16, bold=True, line=1.1) + 0.3)
            rect(s, x + w - tag, my - tth / 2, tag, tth, "head")
            text(s, (x + w - tag + 0.08, my - tth / 2, tag - 0.16, tth), tags[i], tgs, "ink",
                 bold=True, align="c", anchor="m", fit=False, line=1.1)
        if i < n - 1 and (bd or lv.get("tag")):
            hline(s, tx, by + lh + g / 2, x + w - tx, "line", HAIR)


def p_funnel(s, b, p):
    p_pyramid(s, b, p, funnel=True)


def _node(s, x, y, w, h, d, fill, fg, border=None, size=11.5):
    """Org-chart box: bold name over an optional role line."""
    rect(s, x, y, w, h, fill, border)
    name, role = _item(d)
    _pair(s, (x + 0.08, y + 0.04, w - 0.16, h - 0.08), name, role, size, fg,
          fg if fg == "paper" else "text", sub_size=max(size - 2, 8.5))


def _names(d):
    return [t for t in _item(d) if t]


def p_tree(s, b, p):
    """Org chart / 추진체계: root (+ side boxes) → children row → stacked grandchildren, with an
    optional footer strip per child."""
    x, y, w, h = b
    root = p["root"]
    kids = root.get("children", [])
    n = max(len(kids), 1)
    nh = p.get("node_h", min(0.8, max(0.62, h * 0.15)))
    rw, sww, g = min(2.8, w * 0.3), min(2.2, w * 0.2), 0.15
    cw = (w - g * (n - 1)) / n
    sides = [(sd, 1, j) for j, sd in enumerate(root.get("side", []))] + \
            [(sd, -1, j) for j, sd in enumerate(root.get("side_left", []))]
    size = share([(_names(d), (bw - 0.16) * PAIR_W, bh - 0.1) for d, bw, bh in
                  [(root, rw, nh)] + [(k, cw, nh) for k in kids] +
                  [(sd, sww, nh - 0.12) for sd, _, _ in sides]],
                 11.5, max_size=12.5, bold=True, min_size=9, line=1.05, gap=2, underfill=False,
                 name="tree node")
    _node(s, x + (w - rw) / 2, y, rw, nh, root, "ink", "paper", size=size)
    for sd, side, j in sides:  # advisory / PMO boxes beside the root
        edge = x + (w + side * rw) / 2
        sx = edge + 0.5 + j * (sww + 0.15) if side > 0 else edge - 0.5 - sww - j * (sww + 0.15)
        line(s, edge, y + nh / 2, sx if side > 0 else sx + sww, y + nh / 2, "ink2", BORDER,
             dash=True)
        _node(s, sx, y + 0.06, sww, nh - 0.12, sd, "paper", "ink", "line", size=size)
    if not kids:
        return
    bus, ky = y + nh + 0.2, y + nh + 0.4
    cxs = [x + i * (cw + g) + cw / 2 for i in range(n)]
    line(s, x + w / 2, y + nh, x + w / 2, bus, "ink2", BORDER)
    if n > 1:
        line(s, cxs[0], bus, cxs[-1], bus, "ink2", BORDER)
    foot = 0.34 if any(isinstance(k, dict) and k.get("footer") for k in kids) else 0
    ly = ky + nh + 0.12
    avail = y + h - ly - (foot + 0.1 if foot else 0)
    gks = [[] if isinstance(k, str) else k.get("children", []) for k in kids]
    m = max(len(gk) for gk in gks) or 1
    gg = 0.08
    each = min(1.0, (avail - gg * (m - 1)) / m)
    gsize = share([(_names(c), (cw - 0.4) * PAIR_W, each - 0.08) for gk in gks for c in gk], 11,
                  max_size=12, bold=True, min_size=9, line=1.05, gap=2, underfill=False,
                  name="tree leaf")
    for i, (k, gk) in enumerate(zip(kids, gks)):
        kx = x + i * (cw + g)
        hl = p.get("highlight") == i
        line(s, cxs[i], bus, cxs[i], ky, "ink2", BORDER)
        _node(s, kx, ky, cw, nh, k, "accent" if hl else "head", "paper" if hl else "ink",
              None if hl else "line", size=size)
        for j, c in enumerate(gk):
            cy = ly + j * (each + gg)
            line(s, kx + 0.1, cy + each / 2, kx + 0.24, cy + each / 2, "ink2", BORDER)
            _node(s, kx + 0.24, cy, cw - 0.24, each, c, "paper", "ink",
                  "ink" if hl else "line", size=gsize)  # accent stays on the child header only
        if gk:
            line(s, kx + 0.1, ky + nh, kx + 0.1, ly + (len(gk) - 1) * (each + gg) + each / 2,
                 "ink2", BORDER)
        elif isinstance(k, dict) and k.get("items"):
            text(s, (kx + 0.05, ly, cw - 0.1, avail), k["items"], BODY_SIZE, bullets=True,
                 min_size=BODY_MIN, grow=12, spread=True, name=k["name"])
        if isinstance(k, dict) and k.get("footer"):
            rect(s, kx, y + h - foot, cw, foot, "soft", "line", lw=HAIR)
            text(s, (kx + 0.1, y + h - foot, cw - 0.2, foot), k["footer"], 10, "ink", bold=True,
                 align="c", anchor="m", min_size=8.5, name="tree footer")


def p_house(s, b, p):
    """전략 체계도: a label column (비전·목표·추진 전략·추진 기반) beside the roof → goals →
    pillars → base rows."""
    x, y, w, h = b
    lc = p.get("label_width", 1.0)
    x0, cw = x + lc + 0.08, w - lc - 0.08

    def label(ty, th, t):
        if t:
            rect(s, x, ty, lc, th, "ink2")
            text(s, (x + 0.06, ty, lc - 0.12, th), t, 10.5, "paper", bold=True, align="c",
                 anchor="m", min_size=8.5, line=1.1, name="house label")
    rh = 0.8
    label(y, rh, p.get("vision_label", "비전"))
    roof = rect(s, x0, y, cw, rh, "ink", shape=MSO_SHAPE.TRAPEZOID)
    roof.adjustments[0] = cw * 0.08 / rh
    text(s, (x0 + cw * 0.12, y + 0.06, cw * 0.76, rh - 0.12), p["vision"], 14, "paper",
         bold=True, align="c", anchor="m", min_size=10, emph="accent_on_dark", name="vision")
    cy = y + rh + 0.08
    goals = p.get("goals", [])
    if goals:
        gh = 0.5
        label(cy, gh, p.get("goals_label", "목표"))
        gw = (cw - 0.08 * (len(goals) - 1)) / len(goals)
        for i, gl in enumerate(goals):
            gx = x0 + i * (gw + 0.08)
            rect(s, gx, cy, gw, gh, "paper", "line")
            text(s, (gx + 0.1, cy, gw - 0.2, gh), gl, 10.5, "ink", bold=True, align="c",
                 anchor="m", min_size=8.5, line=1.05, name="goal")
        cy += gh + 0.08
    base = p.get("base", [])
    bh = 0.42
    ph = y + h - cy - len(base) * (bh + 0.06)
    pillars = p["pillars"]
    n = len(pillars)
    pg = 0.1
    pw = (cw - pg * (n - 1)) / n
    xs = [x0 + i * (pw + pg) for i in range(n)]
    label(cy, ph, p.get("pillars_label", "추진 전략"))
    hh = 0.44
    foot = 0.38 if any(pl.get("note") for pl in pillars) else 0
    ih = ph - hh - 0.2 - foot
    size = share([(pl["items"], pw - 0.28, ih) for pl in pillars], BODY_SIZE, max_size=12,
                 bullets=True, name="pillars")
    for i, (pl, px) in enumerate(zip(pillars, xs)):
        hl = p.get("highlight") == i
        rect(s, px, cy, pw, hh, "accent" if hl else "head")
        if pl.get("note"):
            rect(s, px, cy + ph - foot, pw, foot, "soft")
            text(s, (px + 0.12, cy + ph - foot, pw - 0.24, foot), "→ " + pl["note"], 10, "ink",
                 bold=True, anchor="m", min_size=8.5, line=1.05, name="pillar note")
        box(s, px, cy, pw, ph, hl)
        text(s, (px + 0.1, cy, pw - 0.2, hh), pl["title"], 11.5, "paper" if hl else "ink",
             bold=True, align="c", anchor="m", min_size=9, line=1.05, name="pillar")
        text(s, (px + 0.14, cy + hh + 0.1, pw - 0.28, ih), pl["items"], size, bullets=True,
             fit=False, spread=True)
    by = cy + ph + 0.06
    if base and all(isinstance(bs, str) for bs in base):  # plain strings share one label cell
        label(by, len(base) * (bh + 0.06) - 0.06, p.get("base_label", "추진 기반"))
    for i, bs in enumerate(base):
        d = bs if isinstance(bs, dict) else {"items": [bs]}
        label(by, bh, d.get("label"))
        its = d["items"]
        m = len(its)
        cells = zip(xs, [pw] * n) if m == n else \
            [(x0 + j * (cw + pg) / m, (cw - pg * (m - 1)) / m) for j in range(m)]
        for (bx, bw), t in zip(cells, its):
            rect(s, bx, by, bw, bh, "head" if i == 0 else "soft")
            text(s, (bx + 0.1, by, bw - 0.2, bh), t, 10.5, "ink" if i == 0 else "text",
                 bold=i == 0, align="c", anchor="m", min_size=8.5, line=1.05, name="base")
        by += bh + 0.06


def _spread_gap(jobs, size, bullets=True, gap=4):
    """One paragraph gap for sibling boxes [(content, w, h)], so their rows line up (spread=True
    would give each box its own)."""
    gs = [(h * 0.9 - text_height(c, size, w, bullets, gap=gap)) * 72 / (len(_paras(c, bullets)) - 1)
          for c, w, h in jobs if c and len(_paras(c, bullets)) > 1]
    return gap + max(0, min(12, min(gs, default=0)))


def p_milestones(s, b, p):
    """연혁·이정표 as a strip of equal cells: a date row over title, text | items[] and an
    optional value at the foot."""
    x, y, w, h = b
    ev = p["events"]
    n = len(ev)
    cw = w / n
    iw = cw - 0.28
    dh = 0.40
    vh = 0.46 if any(e.get("value") for e in ev) else 0
    ts = share([(e["title"], iw, 0.7) for e in ev], 11.5, max_size=12.5, bold=True, min_size=9.5,
               line=1.1, gap=0, underfill=False, name="milestone title")
    th = max(text_height(e["title"], ts, iw, bold=True, line=1.1) for e in ev)
    bt, bb = y + dh + 0.14 + th + 0.1, y + h - vh - 0.1
    body = [e.get("items") or e.get("text") for e in ev]
    size = share([(bd, iw, bb - bt) for bd in body], 10.5, max_size=12, bullets=True,
                 name="milestones")
    gap = _spread_gap([(bd, iw, bb - bt) for bd in body], size)
    for i, (e, bd) in enumerate(zip(ev, body)):
        cx = x + i * cw
        rect(s, cx, y, cw, dh, "accent" if e.get("highlight") else "ink", "paper", lw=1)
        text(s, (cx, y, cw, dh), e["date"], 11, "paper", bold=True, align="c", anchor="m",
             fit=False)
        box(s, cx, y + dh, cw, h - dh)
        text(s, (cx + 0.14, y + dh + 0.14, iw, th + 0.02), e["title"], ts, "ink", bold=True,
             fit=False, line=1.1)
        if bd:
            text(s, (cx + 0.14, bt, iw, bb - bt), bd, size, bullets=isinstance(bd, list),
                 fit=False, gap=gap)
        if e.get("value"):
            hline(s, cx + 0.14, y + h - vh, iw, "line", HAIR)
            text(s, (cx + 0.14, y + h - vh, iw, vh), e["value"], 16, "ink", bold=True,
                 anchor="m", fit=False)
            vw = text_w(e["value"], 16, True) / 0.85 + 0.12
            if e.get("value_label"):
                text(s, (cx + 0.14 + vw, y + h - vh, iw - vw, vh), e["value_label"], 9.5, "muted",
                     anchor="m", line=1.0, min_size=8, name="milestone value label")


def p_profiles(s, b, p):
    x, y, w, h = b
    people = p["people"]
    ncol = p.get("cols", len(people))
    nrow = math.ceil(len(people) / ncol)
    g = 0.12
    cw, chh = (w - g * (ncol - 1)) / ncol, (h - g * (nrow - 1)) / nrow
    pw = min(0.95, cw * 0.28) if any(pr.get("photo") for pr in people) else 0
    metas = [pr.get("meta") or ([f"소속: {pr['org']}"] if pr.get("org") else []) for pr in people]
    kv_h = 0.28 * max(map(len, metas))
    top = 0.56 + max(kv_h, pw * 1.25) + 0.12  # items start here, the same in every card
    foot = 0.32 if any(pr.get("stats") for pr in people) else 0
    bh = chh - top - 0.1 - foot
    if bh < 0.3 and any(pr.get("items") for pr in people):
        warn("profiles: meta rows leave no room for items — cut meta or raise h")
    bh = max(0.2, bh)
    bs = share([(pr.get("items"), cw - 0.28, bh) for pr in people],
               p.get("size", 10.5), 11.5, "profiles", bullets=True)
    for i, (pr, meta) in enumerate(zip(people, metas)):
        cx, cy = x + (i % ncol) * (cw + g), y + (i // ncol) * (chh + g)
        hl = p.get("highlight") == i
        rect(s, cx, cy, cw, 0.44, "accent" if hl else "ink")
        nw = text_w(pr["name"], 12.5, True) + 0.2
        text(s, (cx + 0.12, cy, nw, 0.44), pr["name"], 12.5, "paper", bold=True, anchor="m",
             fit=False, wrap=False)
        text(s, (cx + 0.12 + nw, cy, max(0.4, cw - nw - 0.24), 0.44), pr.get("role", ""), 10,
             "paper" if hl else "on_dark", bold=True, align="r", anchor="m", min_size=8.5,
             line=1.0, name="role")
        kx = cx + 0.12
        src = pr.get("photo") and _src({"src": pr["photo"]})
        if src:
            picture(s, src, kx, cy + 0.56, pw, pw * 1.25, "top")
            rect(s, kx, cy + 0.56, pw, pw * 1.25, None, "line", lw=HAIR)
            kx += pw + 0.14
        kw = cx + cw - 0.12 - kx
        for j, kv in enumerate(meta):  # '라벨: 값' rows
            k, _, v = str(kv).partition(":")
            ky = cy + 0.52 + j * 0.28
            text(s, (kx, ky, 0.6, 0.28), k.strip(), 9, "muted", bold=True, anchor="m", fit=False)
            text(s, (kx + 0.62, ky, max(0.3, kw - 0.62), 0.28), v.strip(), 9.5, "text",
                 anchor="m",
                 min_size=8, line=1.0, name="profile meta")
            if j < len(meta) - 1:
                hline(s, kx, ky + 0.28, kw, "rule", HAIR)
        hline(s, cx + 0.12, cy + top - 0.08, cw - 0.24, "line", HAIR)
        if pr.get("items"):
            _body(s, (cx + 0.14, cy + top, cw - 0.28, bh), pr["items"], bs, True, pr["name"])
        if pr.get("stats"):
            text(s, (cx + 0.14, cy + chh - foot, cw - 0.28, foot - 0.04), pr["stats"], 9,
                 "muted", bold=True, anchor="m", fit=False)
        box(s, cx, cy, cw, chh, hl)


def _lone(body, bullets):
    """A body of one paragraph: the title and body are centred together as one block."""
    return bool(body) and len(_paras(body, bullets)) < 2


def p_numbered(s, b, p):
    x, y, w, h = b
    items = p["items"]
    ncol = p.get("cols", 1)
    per = math.ceil(len(items) / ncol)
    g = 0.3
    cw = (w - g * (ncol - 1)) / ncol
    rh = h / per
    if rh < 0.75:
        warn("numbered: too many items for this row — raise h or use cols:2")
    vc = 1.4 if any(it.get("value") for it in items) else 0
    tw = max(0.5, cw - 0.64 - vc - (0.1 if vc else 0))
    ts = share([(it["title"], tw, 0.34) for it in items], 12.5, 13.5, "numbered title",
               underfill=False, bold=True, min_size=10, line=1.05)
    bodies = [it.get("items") or it.get("text") for it in items]
    bl = any("items" in it for it in items)
    bh = max(0.2, rh - 0.54)
    bs = share([(bd, tw, bh) for bd in bodies], BODY_SIZE, 12, p.get("title", "numbered"),
               bullets=bl, underfill=not all(_lone(bd, bl) for bd in bodies if bd))
    for c in range(ncol):
        hline(s, x + c * (cw + g), y, cw, "ink", FRAME)
    for i, (it, body) in enumerate(zip(items, bodies)):
        cx, cy = x + (i // per) * (cw + g), y + (i % per) * rh
        hl = p.get("highlight") == i
        first, last = i % per == 0, i % per == per - 1 or i == len(items) - 1
        ny = cy + (0 if first else 0.03)  # flush with the list's top and bottom edges
        shape_text(rect(s, cx, ny, 0.5, cy + rh - (0 if last else 0.03) - ny,
                        "accent" if hl else "ink"), i + 1, 14, "paper")
        ty = cy + 0.08
        if _lone(body, "items" in it):
            need = text_height(body, bs, tw, "items" in it)
            ty = cy + max(0.08, (rh - 0.38 - need) / 2)
        text(s, (cx + 0.64, ty, tw, 0.34), it["title"], ts, "ink", bold=True, anchor="m",
             fit=False, line=1.05)
        if _lone(body, "items" in it):
            text(s, (cx + 0.64, ty + 0.38, tw, need + 0.05), body, bs, bullets="items" in it,
                 fit=False, name=it["title"])
        elif body:
            _body(s, (cx + 0.64, cy + 0.46, tw, bh), body, bs, "items" in it, it["title"])
        if it.get("value"):
            vy = cy + rh / 2 - 0.3
            text(s, (cx + cw - vc, vy, vc, 0.36), it["value"], 16, "ink", bold=True, align="r",
                 anchor="m", fit=False, wrap=False)
            text(s, (cx + cw - vc, vy + 0.36, vc, 0.24), it.get("value_label", ""), 8.5, "muted",
                 align="r", min_size=7.5, line=1.0, name="numbered value label")
        if not last:
            hline(s, cx, cy + rh, cw, "line", HAIR)


def _zone(s, x, y, w, h, label, hl=False):
    """Dashed boundary (내부망, 외부 클러스터 …) with a head-filled label tab at its top-left."""
    zr = rect(s, x, y, w, h, "soft", "accent" if hl else "line", lw=1.0 if hl else BORDER)
    zr.line.dash_style = 4
    tw = min(w, text_w(label, 9.5, True) + 0.26)
    rect(s, x, y, tw, 0.28, "accent" if hl else "head")
    text(s, (x + 0.12, y, tw - 0.14, 0.28), label, 9.5, "paper" if hl else "ink", bold=True,
         anchor="m", fit=False)
    return x, y, w, h


def p_flow(s, b, p):
    """Left→right nodes (ink header + bullet body) joined by labelled block arrows; zones
    [{label, from, to}] draw a dashed boundary around a run of nodes."""
    x, y, w, h = b
    nodes, edges, zones = p["nodes"], p.get("edges", []), p.get("zones", [])
    n = len(nodes)
    ag = min(1.2, w * 0.11)
    if zones:  # nodes sit inside the boundaries: inset by the zone padding and tab
        x, w, y, h = x + 0.14, w - 0.28, y + 0.42, h - 0.56
    nw = (w - ag * (n - 1)) / n
    xs = [x + i * (nw + ag) for i in range(n)]
    for z in zones:
        a, e = xs[z["from"]] - 0.14, xs[z["to"]] + nw + 0.14
        _zone(s, a, y - 0.42, e - a, h + 0.56, z["label"], z.get("highlight"))
    hh = 0.42
    bodies = [nd.get("items") or nd.get("text") for nd in nodes]
    ih = h - hh - 0.24
    size = share([(bd, nw - 0.28, ih) for bd in bodies], 10.5, max_size=12, bullets=True,
                 name="flow nodes")
    my = y + h / 2
    aw = min(0.55, ag - 0.3)
    for i, (nd, bd, nx) in enumerate(zip(nodes, bodies, xs)):
        hl = p.get("highlight") == i
        if bd:
            rect(s, nx, y, nw, h, "paper", "accent" if hl else "line", lw=1.25 if hl else BORDER)
        rect(s, nx, y, nw, hh if bd else h, "accent" if hl else "ink")
        text(s, (nx + 0.08, y, nw - 0.16, hh if bd else h), nd["title"], 11.5, "paper",
             bold=True, align="c", anchor="m", min_size=8.5, line=1.05, name="flow node")
        if bd:
            text(s, (nx + 0.14, y + hh + 0.12, nw - 0.28, ih), bd, size,
                 bullets="items" in nd, fit=False, spread=True)
        if i == n - 1:
            continue
        e = edges[i] if i < len(edges) else ""
        fwd, back = (e, None) if isinstance(e, str) else (e.get("label", ""), e.get("back"))
        gx, ax = nx + nw, nx + nw + (ag - aw) / 2
        fy = my - (0.2 if back else 0)
        rect(s, ax, fy - 0.15, aw, 0.30, "ink2", shape=MSO_SHAPE.RIGHT_ARROW)
        if fwd:
            text(s, (gx + 0.04, fy - 0.7, ag - 0.08, 0.5), fwd, 9.5, "ink", bold=True, align="c",
                 anchor="b", line=1.0, min_size=8, name="edge")
        if back:
            rect(s, ax, my + 0.05, aw, 0.30, "grey1", shape=MSO_SHAPE.LEFT_ARROW)
            text(s, (gx + 0.04, my + 0.4, ag - 0.08, 0.5), back, 9, "muted", align="c",
                 line=1.0, min_size=8, name="edge")


def p_progress(s, b, p):
    """Bars that fill the panel: one row per item, optional sub line and target tick."""
    x, y, w, h = b
    items = p["items"]
    mx, n = p.get("max", 100), len(items)
    rh = h / n
    lw = p.get("label_width", w * 0.32)
    shown = [it.get("display", f"{it['value']}%") for it in items]
    vw = min(w * 0.3, max(0.85, max(text_w(d, 12, True) for d in shown) + 0.15))
    tw = w - lw - vw
    bh = min(0.2, rh * 0.35)
    hls = p.get("highlight")
    hls = set(hls) if isinstance(hls, list) else {hls}
    subs = any(it.get("sub") for it in items)
    ls = share([(it["label"], lw - 0.12, rh * (0.5 if subs else 0.9)) for it in items], 11, 12,
               name="progress label", underfill=False, bold=True, min_size=8)
    labelled = False
    for i, it in enumerate(items):
        ry, hl = y + i * rh, i in hls
        if i:
            hline(s, x, ry, w, "line", HAIR)
        if it.get("sub"):
            text(s, (x, ry, lw - 0.12, rh / 2 + 0.02), it["label"], ls, "ink", bold=True,
                 anchor="b", fit=False)
            text(s, (x, ry + rh / 2 + 0.04, lw - 0.12, rh / 2 - 0.06), it["sub"], 9.5, "muted",
                 min_size=8, name="progress sub")
        else:
            text(s, (x, ry, lw - 0.12, rh), it["label"], ls, "ink", bold=True, anchor="m",
                 fit=False)
        by = ry + rh / 2 - bh / 2
        rect(s, x + lw, by, tw, bh, "head")
        rect(s, x + lw, by, tw * min(1, it["value"] / mx), bh, "accent" if hl else "ink2")
        t = it.get("target", p.get("target"))
        if t is not None:
            tx = x + lw + tw * min(1, t / mx)
            line(s, tx, by - 0.06, tx, by + bh + 0.06, "ink", 1.25)
            if rh >= 0.55 and not labelled:  # once: the lead or legend carries the meaning
                text(s, (tx - 0.4, by - 0.25, 0.8, 0.18), p.get("target_label", "목표"), 8, "muted",
                     align="c", anchor="b", fit=False)
                labelled = True
        text(s, (x + w - vw + 0.08, ry, vw - 0.08, rh), shown[i], 12, "accent" if hl else "ink",
             bold=True, align="r", anchor="m", fit=False)


def _chip(it):
    """roadmap chip → (tag | None, text): 'M3 | 문제 5만 건' puts 'M3' in a tag cell."""
    a, sep, z = plain(it).partition(" | ")
    return (a, z) if sep else (None, a)


def p_roadmap(s, b, p):
    """Tracks (rows) × phases (columns) of item chips, stacked and centred in each cell.
    '**chip**' = key item (ink chip); 'M3 | 내용' = month/tag cell + text."""
    x, y, w, h = b
    phases, tracks = p["phases"], p["tracks"]
    lw = p.get("label_width", 1.5)
    pw = (w - lw) / len(phases)
    hh = 0.46
    hl = p.get("highlight")
    th = (h - hh - 0.08) / len(tracks)
    if hl is not None:
        rect(s, x + lw + hl * pw, y + hh, pw, h - hh, "soft")
    for i, ph in enumerate(phases):
        sh = rect(s, x + lw + i * pw, y, pw - 0.02, hh, "accent" if hl == i else "ink2",
                  shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
        sh.adjustments[0] = 0.25
        shape_text(sh, ph, 10.5, "paper", box=(pw - 0.5, hh))
    cells = [[[c for c in ([cell] if isinstance(cell, str) else cell) if c] for cell in t["cells"]]
             for t in tracks]
    cg = 0.06
    tags = [_chip(it)[0] for row in cells for cell in row for it in cell]
    mw = max((text_w(t, 9.5, True) + 0.2 for t in tags if t), default=0)

    def chip_h(k):
        return min(0.62, (th - 0.2 - cg * (k - 1)) / k)
    ks = [max(map(len, row), default=1) or 1 for row in cells]  # one chip grid per track
    size = share([(_chip(it)[1], pw - 0.38 - mw, chip_h(k) - 0.04)
                  for row, k in zip(cells, ks) for cell in row for it in cell], 10, max_size=11,
                 min_size=8.5, line=1.0, gap=0, underfill=False, name="roadmap chips")
    for r, (t, row) in enumerate(zip(tracks, cells)):
        ty = y + hh + 0.08 + r * th
        rect(s, x, ty + 0.04, lw - 0.1, th - 0.08, "ink")
        text(s, (x + 0.08, ty + 0.04, lw - 0.26, th - 0.08), t["name"], 10.5, "paper", bold=True,
             align="c", anchor="m", min_size=8.5, line=1.1, name="track")
        if r < len(tracks) - 1:
            hline(s, x, ty + th, w, "line", HAIR)
        for i, cell in enumerate(row):
            if not cell:
                continue
            ch = chip_h(ks[r])  # row j lines up across phases; short cells end early
            cy0 = ty + (th - ks[r] * ch - cg * (ks[r] - 1)) / 2
            for j, it in enumerate(cell):
                acc = p.get("highlight_chip") == [r, i, j]
                dark = acc or it.startswith("**")
                cx, cy = x + lw + i * pw + 0.1, cy0 + j * (ch + cg)
                rect(s, cx, cy, pw - 0.2, ch, "accent" if acc else "ink" if dark else "paper",
                     None if dark else "line", lw=HAIR)
                tag, body = _chip(it)
                tx = cx + 0.1
                if tag:
                    rect(s, cx, cy, mw, ch, "ink2" if dark else "head")
                    text(s, (cx, cy, mw, ch), tag, 9.5, "paper" if dark else "ink", bold=True,
                         align="c", anchor="m", fit=False)
                    tx = cx + mw + 0.1
                text(s, (tx, cy, cx + pw - 0.28 - tx, ch), body, size,
                     "paper" if dark else "ink", bold=dark, anchor="m", fit=False, line=1.0)


# ------------------------------------------------------------------ icons & images

ICONS_PATH = Path(__file__).resolve().parent.parent / "assets/icons.json"
_ICON = {}


def icon(s, name, x, y, size, color="ink"):
    """Lucide icon (assets/icons.json) rendered to a transparent PNG."""
    import pymupdf
    if "set" not in _ICON:
        _ICON["set"] = json.loads(ICONS_PATH.read_text(encoding="utf-8"))
    icons = _ICON["set"]
    if name not in icons:
        warn(f"unknown icon {name!r} — search the names in assets/icons.json")
        name = "circle-help"
    key = (name, col(color))
    if key not in _ICON:
        svg = ('<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24" '
               f'fill="none" stroke="{col(color)}" stroke-width="2" stroke-linecap="round" '
               f'stroke-linejoin="round">{icons[name]}</svg>')
        pix = pymupdf.open(stream=svg.encode(), filetype="svg")[0].get_pixmap(
            matrix=pymupdf.Matrix(8, 8), alpha=True)
        _ICON[key] = pix.tobytes("png")
    return s.shapes.add_picture(io.BytesIO(_ICON[key]), Inches(x), Inches(y), Inches(size),
                                Inches(size))


def alpha(sh, a):
    """Make a solid-filled shape translucent (a = opacity 0..1)."""
    clr = sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha"), val=str(int(a * 100000)))
    return sh


def _src(p):
    src = p.get("src") and (CTX["base"] / p["src"])
    return src if src and src.exists() else None


def aspect(src, region=None):
    """Pixel aspect (w/h) of an image file, or of a region [fx, fy, fw, fh] of it."""
    from pptx.parts.image import Image
    iw, ih = Image.from_file(str(src)).size
    fw, fh = (region[2], region[3]) if region else (1, 1)
    return fw * iw / (fh * ih)


def picture(s, src, x, y, w, h, crop="center", focus=None):
    """Fill (x, y, w, h) with the image, cropping the overflow.
    crop: center | top (keeps the top-left corner — screenshots).
    focus=[fx, fy, fw, fh] (fractions of the source) aims the crop at a region; the region is
    widened to the box shape, never distorted. Returns the shown region (fx, fy, fw, fh) so
    callers can map source points onto the slide with map_pt()."""
    pic = s.shapes.add_picture(str(src), Inches(x), Inches(y), Inches(w), Inches(h))
    iw, ih = pic.image.size
    ba = w / h
    if focus:
        fx, fy, fw0, fh0 = focus
        cx, cy = fx + fw0 / 2, fy + fh0 / 2
        fw, fh = fw0, fh0
        if fw * iw / (fh * ih) < ba:
            fw = fh * ih * ba / iw
        else:
            fh = fw * iw / (ba * ih)
        if fw > 1:
            fw, fh = 1, iw / (ba * ih)
        if fh > 1:
            fh, fw = 1, ih * ba / iw
        if max(fw / fw0, fh / fh0) > 1.4:
            warn(f"focus {focus} widened {max(fw / fw0, fh / fh0):.1f}x to fit its box — "
                 "pick a region closer to the box shape")
        fx, fy = min(max(cx - fw / 2, 0), 1 - fw), min(max(cy - fh / 2, 0), 1 - fh)
    elif iw / ih > ba:  # too wide → trim sides (top crop keeps the left edge)
        fw, fh = ba / (iw / ih), 1
        fx, fy = (0 if crop == "top" else (1 - fw) / 2), 0
    else:  # too tall → trim bottom (screenshots) or both ends (photos)
        fw, fh = 1, (iw / ih) / ba
        fx, fy = 0, (0 if crop == "top" else (1 - fh) / 2)
    pic.crop_left, pic.crop_top = max(fx, 0), max(fy, 0)
    pic.crop_right, pic.crop_bottom = max(1 - fx - fw, 0), max(1 - fy - fh, 0)
    return fx, fy, fw, fh


def map_pt(scr, shown, px, py):
    """Source-image point (fractions) → slide inches, for a picture drawn into scr."""
    (sx, sy, sw, sh), (fx, fy, fw, fh) = scr, shown
    return sx + (px - fx) / fw * sw, sy + (py - fy) / fh * sh


def caption(s, x, y, w, cap, fig=True):
    """Figure caption strip under an image: '[그림 n] caption' (deck-wide numbering)."""
    rect(s, x, y, w, CAP_H, "soft")
    if fig:
        CTX["fig"] += 1
        cap = f"**[그림 {CTX['fig']}]** {cap}"
    text(s, (x + 0.1, y, w - 0.2, CAP_H), cap, 9.5, "text", anchor="m", emph="ink", min_size=8,
         line=1.0, name="caption")


def placeholder(s, x, y, w, h, alt, dark=False):
    rect(s, x, y, w, h, "ink2" if dark else "soft", None if dark else "line")
    d = min(0.4, w * 0.2, h * 0.3)
    icon(s, "image", x + w / 2 - d / 2, y + h / 2 - d * 0.9, d, "on_dark" if dark else "grey1")
    text(s, (x + 0.1, y + h / 2 + d * 0.25, w - 0.2, 0.5), alt or "이미지를 넣어 주세요", 9,
         "on_dark" if dark else "muted", align="c", fit=False)


FRAME_RATIO = {"browser": None, "laptop": 16 / 10, "tablet": 4 / 3, "phone": 9 / 19.5}
STRIP = 0.26  # browser title strip


def _img_fit(b, a, chrome=0.0, tol=1.3):
    """Shrink box b so the picture under its chrome strip is within tol of aspect a (the focus
    widens by up to tol to fill the rest). Centred across, top-aligned."""
    x, y, w, h = b
    ba = w / max(h - chrome, 0.01)
    if ba > a * tol:
        nw = (h - chrome) * a * tol
        return x + (w - nw) / 2, y, nw, h
    if ba < a / tol:
        return x, y, w, w * tol / a + chrome
    return b


def _frame(s, b, kind, src, alt, focus=None, label=None, cell=True):
    """Device frame around a screenshot; returns (screen rect, shown source region, outer rect).
    browser: an ink2 title strip on a bordered page; None: the bare bordered screenshot.
    laptop / tablet / phone: drawn at 92% on a soft figure cell that fills b, or with cell=False
    as large as fits, bottom-centred in b (device compositions, split mats)."""
    x, y, w, h = b
    out, k = b, 0.92 if cell else 1
    if kind in ("browser", None):
        rect(s, x, y, w, h, "paper", "line", lw=BORDER)
        top = STRIP if kind else 0
        if kind:
            rect(s, x, y, w, top, "ink2")
            text(s, (x + 0.1, y, w - 0.2, top), "■ " + (label or alt or ""), 8.5, "paper",
                 anchor="m", fit=False, wrap=False)
        scr = (x, y + top, w, h - top)
    else:
        if cell:
            rect(s, x, y, w, h, "soft")
        r = FRAME_RATIO[kind]
        if kind == "laptop":  # ink body, 2.5% bezel, flat grey2 base
            sw = min(w * (0.94 if cell else 1 / 1.1), (h * k - 0.09) * r)
            sh = sw / r
            sx = x + (w - sw) / 2
            sy = y + (h - sh - 0.09) / 2 if cell else y + h - 0.09 - sh
            rect(s, sx, sy, sw, sh, "ink")
            bw = min(1.1 * sw, w)  # flat base, kept on the figure cell
            rect(s, x + (w - bw) / 2, sy + sh + 0.02, bw, 0.07, "grey2")
            bz = sw * 0.025
            scr, out = (sx + bz, sy + bz, sw - 2 * bz, sh - 2 * bz), (x + (w - bw) / 2, sy, bw, sh + 0.09)
        else:
            fh = min(h * k, w * k / r)
            fw = fh * r
            fx, fy = x + (w - fw) / 2, y + (h - fh) / 2 if cell else y + h - fh
            body = rect(s, fx, fy, fw, fh, "ink", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            body.adjustments[0] = 0.10 if kind == "phone" else 0.05
            bz = fw * (0.045 if kind == "phone" else 0.05)
            top = bz * (2 if kind == "phone" else 1)
            scr, out = (fx + bz, fy + top, fw - 2 * bz, fh - 2 * top), (fx, fy, fw, fh)
    if src:
        k = aspect(src, focus) * scr[3] / scr[2] if focus and kind not in ("browser", None) else 1
        if 1.15 < max(k, 1 / k) <= 1.4:  # picture() warns above 1.4
            warn(f"{kind} screen is {scr[2] / scr[3]:.2f}:1, so focus {focus} is widened "
                 f"{max(k, 1 / k):.2f}x — pick a region near that shape or drop focus")
        return scr, picture(s, src, *scr, crop="top", focus=focus), out
    placeholder(s, *scr, alt, dark=kind not in ("browser", None))
    return scr, (0, 0, 1, 1), out


def _img_mark(s, x, y, n, d=0.28):
    """Square ink marker centred on (x, y): paper outline, number n."""
    shape_text(rect(s, x - d / 2, y - d / 2, d, d, "ink", "paper", lw=1.5), n, 10, "paper")


def _img_clamp(scr, mx, my, d=0.14):
    """Keep a marker centre far enough inside the screen that the whole marker stays on it."""
    sx, sy, sw, sh = scr
    return min(max(mx, sx + d), sx + sw - d), min(max(my, sy + d), sy + sh - d)


def _img_corner(scr, x1, y1, x2, y2, d=0.14):
    """Marker spot just outside a region: its top-right corner, else the first of left / right
    of its middle, top-left, bottom-right, bottom-left that keeps the marker on the screen, else
    top-right clamped onto it."""
    ym = (y1 + y2) / 2
    spots = [(x2 + d, y1 - d), (x1 - d, ym), (x2 + d, ym), (x1 - d, y1 - d), (x2 + d, y2 + d),
             (x1 - d, y2 + d)]
    on = [m for m in spots if math.dist(_img_clamp(scr, *m), m) < 0.03]
    return on[0] if on else _img_clamp(scr, *spots[0])


def _img_parts(t):
    """'title\\nline\\nline' → (title, [lines])."""
    t = str(t).split("\n")
    return t[0], t[1:]


def _img_rows(s, x, ys, w, rows, nums, hs, name):
    """Numbered spec rows at tops ys: an ink No cell, a bold title and ○ body lines; one shared
    title size and one body size."""
    tw = w - 0.66
    parts = [_img_parts(r) for r in rows]
    ts = share([(t, tw, 0.34) for t, _ in parts], 11.5, 12.5, name, underfill=False, bold=True,
               min_size=10)
    bs = share([(b, tw, rh - 0.56) for (_, b), rh in zip(parts, hs)], 10.5, 12, name, bullets=True)
    for (t, b), n, y, rh in zip(parts, nums, ys, hs):
        shape_text(rect(s, x, y + 0.04, 0.42, rh - 0.08, "ink"), n, 11, "paper")
        text(s, (x + 0.54, y + 0.08, tw, 0.34), t, ts, "ink", bold=True, anchor="m", fit=False)
        if b:
            text(s, (x + 0.54, y + 0.46, tw, rh - 0.56), b, bs, bullets=True, fit=False,
                 spread=True)


def _img_legend(s, x, y, w, h, rows, nums, title):
    """Spec legend: a No | title header, then one row per callout down the full height.
    Returns the row centres (for leaders)."""
    rect(s, x, y, w, 0.32, "head")
    hline(s, x, y, w, "ink", FRAME)
    text(s, (x, y, 0.42, 0.32), "No", 10, "ink", bold=True, align="c", anchor="m", fit=False)
    text(s, (x + 0.54, y, w - 0.64, 0.32), title, 10, "ink", bold=True, anchor="m", fit=False)
    rh = (h - 0.32) / len(rows)
    _img_rows(s, x, [y + 0.32 + i * rh for i in range(len(rows))], w, [c["text"] for c in rows],
              nums, [rh] * len(rows), "legend")
    for i in range(1, len(rows)):
        hline(s, x, y + 0.32 + i * rh, w, "line", HAIR)
    hline(s, x, y + h, w, "ink", BORDER)
    return [y + 0.32 + (i + 0.5) * rh for i in range(len(rows))]


def _img_zooms(s, x, y, w, h, zooms, n0, src, scr, shown):
    """Magnified insets of source regions stacked in a column, joined to their regions."""
    n, g = len(zooms), 0.16
    ras = [aspect(src, z["region"]) for z in zooms]
    nat = [w / a for a in ras]
    need = max(0.62 + text_height(_img_parts(z.get("text", ""))[1], 10.5, w - 0.66, True)
               for z in zooms)
    top = max(nat)  # cap the tallest insets until the text rows fit. ponytail: linear search
    while top > 0.5 and sum(min(v, top) for v in nat) > h - g * (n - 1) - need * n:
        top -= 0.02
    ihs = [min(v, top) for v in nat]
    th = (h - g * (n - 1) - sum(ihs)) / n
    sx, sy, sw, sh = shown
    cy, ys = y, []
    for i, (z, ih, a) in enumerate(zip(zooms, ihs, ras)):
        c = "accent" if z.get("highlight") else "ink"
        iw = min(w, ih * a * 1.15)  # a capped inset narrows rather than drag in neighbours
        fx, fy, fw, fh = picture(s, src, x, cy, iw, ih, focus=z["region"])  # region as shown
        rect(s, x, cy, iw, ih, None, c, lw=FRAME)
        x0, y0 = max(fx, sx), max(fy, sy)  # the part of it the main picture shows
        x1, y1 = min(fx + fw, sx + sw), min(fy + fh, sy + sh)
        if x1 <= x0 or y1 <= y0:
            warn(f"zoom region {z['region']} is outside the visible image — move it or widen "
                 "the focus")
        else:
            (rx, ry), (rx2, ry2) = map_pt(scr, shown, x0, y0), map_pt(scr, shown, x1, y1)
            rect(s, rx, ry, rx2 - rx, ry2 - ry, None, c, lw=1.25)
            line(s, rx2, ry, x, cy, c, HAIR)
            line(s, rx2, ry2, x, cy + ih, c, HAIR)
            _img_mark(s, *_img_corner(scr, rx, ry, rx2, ry2), n0 + i)
        ys.append(cy + ih + 0.06)
        cy += ih + th + g
    _img_rows(s, x, ys, w, [z.get("text", "") for z in zooms], range(n0, n0 + n), [th - 0.06] * n,
              "zoom")


def _img_labels(s, sides, cols, lw, y, h, scr, shown):
    """Leader-line annotations in a column on each side of the image, each label kept level with
    its point where it fits. sides = {"l": [...], "r": [...]}; cols = {"l": x, "r": x}."""
    labs = [(k, c) for k in sides for c in sides[k]]
    parts = [_img_parts(c["text"]) for _, c in labs]
    ts = share([(t, lw, 0.34) for t, _ in parts], 11.5, 12.5, "label", underfill=False,
               bold=True, min_size=10)
    bs = share([(b, lw, h / len(sides[k]) - 0.56) for (k, _), (_, b) in zip(labs, parts)], 10.5,
               12, "label", bullets=True)
    for k in sides:
        lx, left = cols[k], k == "l"
        own = [(map_pt(scr, shown, c["x"], c["y"]), _img_parts(c["text"])) for kk, c in labs
               if kk == k]
        own.sort(key=lambda o: o[0][1])
        need = [0.56 + text_height(b, bs, lw, True) for _, (_, b) in own]
        tops, prev = [], y
        for ((_, py), _), nd in zip(own, need):  # level with its point, below the one above …
            tops.append(max(py, prev))
            prev = tops[-1] + nd
        nxt = y + h
        for i in reversed(range(len(own))):  # … then back up to stay inside the panel
            tops[i] = nxt = max(y, min(tops[i], nxt - need[i]))
        if len(own) > 1 and tops[-1] + need[-1] - tops[0] < h * 0.85:  # ragged: spread evenly
            g = (h - sum(need)) / (len(own) - 1)
            tops = [y + sum(need[:i]) + g * i for i in range(len(own))]
        edge = scr[0] - 0.22 if left else scr[0] + scr[2] + 0.22
        inner = lx + lw if left else lx
        for ((px, py), (tt, b)), t, nd in zip(own, tops, need):
            hline(s, lx, t, lw, "ink", FRAME)
            line(s, inner, t, edge, t, "ink2", BORDER)
            line(s, edge, t, px, py, "ink2", BORDER)
            rect(s, px - 0.06, py - 0.06, 0.12, 0.12, "ink", "paper", lw=1)
            text(s, (lx, t + 0.06, lw, 0.34), tt, ts, "ink", bold=True, anchor="m", fit=False)
            if b:
                text(s, (lx, t + 0.46, lw, nd - 0.5), b, bs, bullets=True, fit=False)


def _img_stats(s, x, y, w, stats, hl):
    """Key-figure cards (value over label) sitting across the image's foot."""
    n, g, ch = len(stats), 0.16, 0.9
    cw = (w - 0.6 - g * (n - 1)) / n
    vs = 24
    while vs > 14 and any(text_w(str(st["value"]), vs, True) > (cw - 0.3) * 0.9 for st in stats):
        vs -= 1
    ls = share([(st.get("label", ""), cw - 0.28, ch - 0.56) for st in stats], 10, 10.5,
               "stat label", underfill=False, min_size=8.5, line=1.05)
    for i, st in enumerate(stats):
        cx, on = x + 0.3 + i * (cw + g), hl == i
        box(s, cx, y, cw, ch, on, "paper")
        text(s, (cx + 0.14, y + 0.08, cw - 0.28, 0.44), str(st["value"]), vs,
             "accent" if on else "ink", bold=True, anchor="m", fit=False)
        text(s, (cx + 0.14, y + 0.52, cw - 0.28, ch - 0.58), st.get("label", ""), ls, fit=False,
             line=1.05)


def p_image(s, b, p):
    """Screenshot / photo, optionally in a device frame, with numbered callouts + spec legend,
    zoom insets, leader-line labels, stat cards across its foot and a figure caption.
    Every x / y / region is a fraction of the source image, so it survives any focus crop."""
    x, y, w, h = b
    src, frame, focus = _src(p), p.get("frame"), p.get("focus")
    calls = p.get("callouts", [])
    rows = [c for c in calls if c.get("text")]
    zooms = p.get("zooms", []) if src else []
    stats, cap, bleed = p.get("stats", []), p.get("caption"), p.get("bleed")
    sides = {"l": [], "r": []}
    for c in p.get("labels", []):  # side: l | r (left / right also work)
        sides["r" if str(c.get("side") or "lr"[c["x"] >= 0.5])[0] == "r" else "l"].append(c)
    lw = w * p.get("legend_width", 0.36) if rows or zooms else 0
    llw = w * p.get("label_width", 0.28)
    ax = x + (llw + 0.3 if sides["l"] else 0)
    aw = x + w - (lw + 0.24 if lw else 0) - (llw + 0.3 if sides["r"] else 0) - ax
    ah = h - (CAP_H + 0.08 if cap else 0) - (0.45 if stats else 0)
    fb = (ax, y, aw, ah)
    if bleed:  # run past the slide margin(s) the panel touches
        x0 = 0 if abs(ax - M) < 0.01 else ax
        x1 = W if abs(ax + aw - (W - M)) < 0.01 else ax + aw
        fb = (x0, y, x1 - x0, ah)
    elif src and frame in (None, "browser"):
        fb = _img_fit(fb, aspect(src, focus), STRIP if frame else 0,
                      1.0 if p.get("fit") == "contain" else 1.3)
        if fb[2] * fb[3] < aw * ah * 0.8:
            warn(f"image fills only {fb[2] * fb[3] / (aw * ah):.0%} of its box — pick a focus "
                 f"near {aw / (ah - (STRIP if frame else 0)):.1f}:1 or change the row h")
    if frame:
        scr, shown, _ = _frame(s, fb, frame, src, p.get("alt"), focus, p.get("screen"))
    elif src:
        shown, scr = picture(s, src, *fb, p.get("crop", "center"), focus), fb
        if not bleed and p.get("border", True):
            rect(s, *fb, None, "line", lw=BORDER)
    else:
        placeholder(s, *fb, p.get("alt"))
        scr, shown = fb, (0, 0, 1, 1)
    marks, starts = [], []  # marker centres · where each leader leaves (a region's right edge)
    for c in calls:
        mx, my = map_pt(scr, shown, c["x"], c["y"])
        if c.get("box"):  # region: [w, h] in source fractions, centred on x, y
            bw, bh = c["box"][0] / shown[2] * scr[2], c["box"][1] / shown[3] * scr[3]
            rect(s, mx - bw / 2, my - bh / 2, bw, bh, None, "accent", lw=2)
            starts.append((mx + bw / 2, my))
            mx, my = _img_corner(scr, mx - bw / 2, my - bh / 2, mx + bw / 2, my + bh / 2)
        else:
            mx, my = _img_clamp(scr, mx, my)
            starts.append((mx + 0.14, my))
        marks.append((mx, my))
    lx = x + w - lw
    zh = (h if not rows else h * 0.5) if zooms else 0
    if zooms:
        _img_zooms(s, lx, y, lw, zh - (0.16 if rows else 0), zooms, len(calls) + 1, src, scr,
                   shown)
    if rows:
        mids = _img_legend(s, lx, y + zh, lw, h - zh, rows, [calls.index(c) + 1 for c in rows],
                           p.get("legend_title", "화면 구성"))
        if p.get("leaders"):  # elbow from each marker to its row
            xr = fb[0] + fb[2]
            for j, (c, ym) in enumerate(zip(rows, mids)):
                mx, my = starts[calls.index(c)]
                xa = xr + 0.05 + j * min(0.04, 0.14 / len(rows))
                pts = [(mx, my), (xa, my), (xa, ym), (lx, ym)]
                for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
                    line(s, x1, y1, x2, y2, "ink2", BORDER)
    for i, (mx, my) in enumerate(marks):
        _img_mark(s, mx, my, i + 1)
    sides = {k: v for k, v in sides.items() if v}
    if sides:
        _img_labels(s, sides, {"l": x, "r": x + w - lw - llw - (0.24 if lw else 0)}, llw, y, h,
                    scr, shown)
    cx0, cx1 = max(fb[0], M), min(fb[0] + fb[2], W - M)
    if stats:
        _img_stats(s, cx0, fb[1] + fb[3] - 0.45, cx1 - cx0, stats, p.get("highlight"))
    if cap:
        caption(s, cx0, fb[1] + fb[3] + (0.45 if stats else 0) + 0.08, cx1 - cx0, cap,
                p.get("fig", True))


def _img_tag(it, i, p):
    return it.get("tag", f"{i + 1}단계" if p.get("layout") == "flow" else None)


def _img_band(cap, tag, big, cw):
    """Caption size, pad and title-line height for an ink band (caption fitted to one line)."""
    pad, cs = (0.12, 12.5) if big else (0.08, 11)
    room = cw - 0.28 - (text_w(tag, 8.5, True) + 0.34 if tag else 0)
    while cap and cs > 9 and text_w(cap, cs, True) > room:
        cs -= 0.5
    if cap and text_w(cap, cs, True) > room:
        warn(f"image caption {cap!r} does not fit its band on one line — shorten it")
    return cs, pad, max(0.28, cs * 1.35 / 72)


def _img_tile(s, c, it, i, p, over, band, ih, ts):
    """One image of an images panel: picture, then either a tag cell + caption strip + text box
    under it, or ("overlay", band = _img_band()) a solid ink band under it holding tag, caption
    and text."""
    cx, cy, cw, ch = c
    src, cap, txt = _src(it), it.get("caption"), it.get("text")
    tag = _img_tag(it, i, p)
    hl = p.get("highlight") == i
    tw = text_w(tag, 8.5, True) + 0.24 if tag else 0
    if band:  # size the band first, so the picture's focus shows whole above it
        cs, pad, lh = band
        th = text_height(txt, ts, cw - 0.28) + 0.04 if txt else 0
        bh = min(lh + 2 * pad + th, ch * 0.5)
    strip = CAP_H if not over and (cap or tag) else 0
    ih = ih or ch - strip - (bh if band else 0)
    if src:
        picture(s, src, cx, cy, cw, ih, it.get("crop", p.get("crop", "center")), it.get("focus"))
        rect(s, cx, cy, cw, ih, None, "line", lw=BORDER)
    else:
        placeholder(s, cx, cy, cw, ih, it.get("alt"))
    if band:
        by = cy + ih
        rect(s, cx, by, cw, bh, "ink")
        if tag:
            shape_text(rect(s, cx + 0.14, by + pad + (lh - 0.24) / 2, tw, 0.24,
                            "accent" if hl else "paper"), tag, 8.5, "paper" if hl else "ink")
        if cap:
            text(s, (cx + 0.14 + (tw + 0.1 if tag else 0), by + pad, cw - 0.28 - tw, lh), cap, cs,
                 "paper", bold=True, anchor="m", fit=False, wrap=False)
        if txt:
            text(s, (cx + 0.14, by + pad + lh, cw - 0.28, th), txt, ts, "on_dark", fit=False)
        return
    if strip:
        if tag:
            shape_text(rect(s, cx, cy + ih, tw, CAP_H, "accent" if hl else "ink"), tag, 8.5,
                       "paper")
        if cap:
            caption(s, cx + tw, cy + ih, cw - tw, cap, p.get("fig", True))
        else:
            rect(s, cx + tw, cy + ih, cw - tw, CAP_H, "soft")
    if txt:
        ty = cy + ih + strip + 0.08
        box(s, cx, ty, cw, cy + ch - ty)
        text(s, (cx + 0.14, ty + 0.1, cw - 0.28, cy + ch - ty - 0.2), txt, ts, bullets=True,
             fit=False, spread=True)


def p_images(s, b, p):
    """Several screenshots / photos. layout: grid · feature / mosaic (one hero + a column of the
    rest; hero_side left | right) · flow (one row joined by arrows). Captions go in a strip under
    each image, or in a solid ink band under it ("overlay", the mosaic default); item text goes
    in a box sized to it, the images take the rest of the height."""
    x, y, w, h = b
    items, lay = p["items"], p.get("layout", "grid")
    n = len(items)
    over = p.get("caption_style", "overlay" if lay == "mosaic" else "strip") == "overlay"
    g = p.get("gap", {"mosaic": 0.08, "flow": 0.45}.get(lay, 0.14))
    hero = lay in ("feature", "mosaic") and n > 1
    k = p.get("feature", 0) if hero else None
    if hero:
        hw = w * p.get("hero_w", 0.6) - g / 2
        rest = [i for i in range(n) if i != k]
        rc = 1 if len(rest) <= 3 else 2
        rr = math.ceil(len(rest) / rc)
        rw, rh = (w - hw - g * rc) / rc, (h - g * (rr - 1)) / rr
        hx, ox = (x + w - hw, x) if p.get("hero_side") == "right" else (x, x + hw + g)
        cells = {k: (hx, y, hw, h)}
        for j, i in enumerate(rest):
            cells[i] = (ox + (j % rc) * (rw + g), y + (j // rc) * (rh + g), rw, rh)
    else:
        nc = n if lay == "flow" else p.get("cols", min(n, 3))
        cw, ch = (w - g * (nc - 1)) / nc, (h - g * (math.ceil(n / nc) - 1)) / math.ceil(n / nc)
        cells = {i: (x + (i % nc) * (cw + g), y + (i // nc) * (ch + g), cw, ch) for i in range(n)}
    ihs, ts, bands = {}, 10.5, {}
    txts = [i for i in range(n) if items[i].get("text")]
    if over:
        bands = {i: _img_band(it.get("caption"), _img_tag(it, i, p), i == k, cells[i][2])
                 for i, it in enumerate(items)
                 if it.get("caption") or _img_tag(it, i, p) or it.get("text")}
    if over and txts:  # band text: one size for every tile, at most half a tile tall
        ts = share([(items[i]["text"], cells[i][2] - 0.28,
                     cells[i][3] * 0.5 - bands[i][2] - 2 * bands[i][1] - 0.04) for i in txts],
                   10.5, 11.5, "image band text", underfill=False)
    elif txts:  # text box sized to its text; the images (row siblings share a height) get the rest
        asp = [aspect(_src(it), it.get("focus")) if _src(it) else 1.6 for it in items]
        for idx in ([[i] for i in txts] if hero else [list(range(n))]):
            cx, cy, cw, ch = cells[idx[0]]
            lo, hi = min(asp[j] for j in idx), max(asp[j] for j in idx)
            need = max(text_height(items[j].get("text", ""), 10.5, cw - 0.28, True) for j in idx)
            room = ch - CAP_H - 0.08 - max(0.5, need + 0.3)
            ih = min(room, cw / hi * 1.3)
            if ih < cw / lo / 1.3:
                warn("images row too short for image + text — raise the row 'h' or shorten "
                     "the text")
                ih = max(0.6, ih)
            elif room - ih > 0.25:
                warn(f"images row is {room - ih:.1f}in taller than images + text need — lower "
                     "the row 'h', add text or pick taller focus regions")
            ihs.update((j, ih) for j in idx)
        ts = share([(items[i]["text"], cells[i][2] - 0.28, cells[i][3] - ihs[i] - CAP_H - 0.28)
                    for i in txts], 10.5, 11.5, "image text", bullets=True)
    for i, it in enumerate(items):
        _img_tile(s, cells[i], it, i, p, over, bands.get(i), ihs.get(i), ts)
    if lay == "flow":
        for i in range(n - 1):
            cx, cy, cw, ch = cells[i]
            my = cy + (ihs.get(i) or ch - CAP_H) / 2
            rect(s, cx + cw + g / 2 - 0.15, my - 0.18, 0.30, 0.36, "ink2",
                 shape=MSO_SHAPE.RIGHT_ARROW)


# ------------------------------------------------------------------ devices, storyboard, compare

def p_devices(s, b, p):
    """Laptop + phone overlapping on one figure cell: the same service on web and mobile."""
    x, y, w, h = b
    lap, ph = p.get("laptop") or {}, p.get("phone") or {}
    cap = p.get("caption")
    fh = h - (CAP_H if cap else 0)
    box(s, x, y, w, fh, fill="soft")
    floor = 0.36 if lap.get("label") or ph.get("label") else 0.14
    aw, ah = w - 0.5, fh - 0.3 - floor
    # laptop body bw wide (+10% base), phone 85% of the laptop height, overlapping 40% of its width;
    # a lone device takes the whole figure
    bw = min((aw - 0.02) / (1.247 if ph else 1.1), (ah - 0.09) * 1.6) if lap else 0
    lh = bw / 1.6 + 0.09
    phh = 0.85 * lh if lap else min(ah, aw * 19.5 / 9)
    phw = phh * 9 / 19.5
    tot = (1.1 * bw if lap else 0) + ((0.6 if lap else 1) * phw if ph else 0)
    gx, gb = x + 0.25 + (aw - tot) / 2, y + 0.3 + ah
    left = p.get("phone_side") == "left"
    lx, px = (gx + tot - 1.1 * bw, gx) if left else (gx, gx + tot - phw)
    lo = lap and _frame(s, (lx, gb - lh, 1.1 * bw, lh), "laptop", _src(lap), lap.get("alt"),
                        lap.get("focus"), cell=False)[2]
    po = ph and _frame(s, (px, gb - phh, phw, phh), "phone", _src(ph), ph.get("alt"),
                       ph.get("focus"), cell=False)[2]
    for o, d, al in ((lo, lap, "r" if left else "l"), (po, ph, "l" if left else "r")):
        if o and d.get("label"):
            lw_ = text_w(d["label"], 9.5, True) * 1.15 + 0.1
            tx = o[0] + o[2] - lw_ if al == "r" else o[0]
            text(s, (tx, gb + 0.08, lw_, 0.24), d["label"], 9.5, "ink", bold=True, align=al,
                 anchor="m", fit=False)
    if cap:
        caption(s, x, y + fh, w, cap, p.get("fig", True))


def p_storyboard(s, b, p):
    """Demo scenario on a rail: numbered step headers, then per step a screen, bullets and a metric
    strip; screens, texts and metrics share one row each."""
    x, y, w, h = b
    items, hl = p.get("items", []), p.get("highlight")
    if not items:
        return
    n, g, rail = len(items), 0.22, 0.34
    cw = (w - g * (n - 1)) / n
    mh = 0.40 if any(it.get("metric") for it in items) else 0
    texts = [it.get("text") or [] for it in items]
    free = h - rail - 0.36 - mh  # screen + text
    need = max((text_height(t, 11, cw, True) for t in texts if t), default=0)
    # text grows only while no step wraps an extra line, unless the screens are already capped
    tb = max(need * 12.5 / 11 / 0.9, free - cw)
    size = share([(t, cw, tb) for t in texts], 11, 12.5, "storyboard text", underfill=False,
                 bullets=True)
    need = max((text_height(t, size, cw, True) for t in texts if t), default=0)
    ih = min(cw, free - need)  # screens take what the text leaves
    if ih < 0.8:
        warn("storyboard row too short for screens — raise its h or cut text")
        ih = 0.8
    ty = y + rail + 0.12 + ih + 0.14
    th = max(need, y + h - mh - 0.10 - ty)
    k = max(len(_paras(t, True)) for t in texts)  # one gap for all steps keeps their rhythm
    gap = 4 + (min(12, (th * 0.9 - need) * 72 / (k - 1)) if k > 1 and need < th * 0.85 else 0)
    hline(s, x, y + rail / 2, w, "ink2", 1.25)  # the rail; step cells and tags sit on it
    for i, it in enumerate(items):
        cx, on = x + i * (cw + g), hl == i
        shape_text(rect(s, cx, y, 0.34, rail, "accent" if on else "ink"), i + 1, 11, "paper")
        tm = it.get("time")
        tmw = text_w(tm, 9, True) * 1.15 + 0.2 if tm else 0
        if tm:
            shape_text(rect(s, cx + cw - tmw, y + 0.04, tmw, rail - 0.08, "paper", "ink2", lw=HAIR),
                       tm, 9, "ink2")
        t = it.get("title", "")
        tw_ = min(text_w(t, 11.5, True) + 0.22, cw - 0.34 - tmw - 0.12)
        rect(s, cx + 0.34, y, tw_, rail, "paper")  # knock-out: the title cuts the rail
        text(s, (cx + 0.44, y, tw_ - 0.1, rail), t, 11.5, "ink", bold=True, anchor="m", min_size=9,
             name="storyboard title")
        sy, src = y + rail + 0.12, _src(it)
        if src:
            picture(s, src, cx, sy, cw, ih, "top", focus=it.get("focus"))
            rect(s, cx, sy, cw, ih, None, "accent" if on else "line", lw=FRAME if on else BORDER)
        else:
            placeholder(s, cx, sy, cw, ih, it.get("alt"))
        text(s, (cx, ty, cw, th), texts[i], size, bullets=True, fit=False, gap=gap)
        if it.get("metric"):
            rect(s, cx, y + h - mh, cw, mh, "accent_soft" if on else "head")
            text(s, (cx + 0.12, y + h - mh, cw - 0.24, mh), it["metric"], 10.5,
                 "accent" if on else "ink", bold=True, anchor="m", min_size=8.5, emph="ink",
                 name="storyboard metric")


def p_compare(s, b, p):
    """Before / after: a header band, a screen and paired metric rows per side, with per-row
    changes in the gutter and an optional delta badge straddling both screens."""
    x, y, w, h = b
    rows, gut, hh = p.get("rows", []), 0.8, 0.36
    cw = (w - gut) / 2
    rh = min(0.52, max(0.36, (h - hh - 0.1) * 0.4 / max(len(rows), 1))) if rows else 0
    ih = h - hh - (0.1 + len(rows) * rh if rows else 0)
    for k, key in enumerate(("before", "after")):
        sd, cx, aft = p[key], x + k * (cw + gut), k == 1
        rect(s, cx, y, cw, hh, "ink" if aft else "head")
        text(s, (cx + 0.12, y, cw - 0.24, hh), sd.get("title", ""), 11.5, "paper" if aft else "ink",
             bold=True, anchor="m", min_size=9, name="compare title")
        src = _src(sd)
        if src:
            picture(s, src, cx, y + hh, cw, ih, "top", focus=sd.get("focus"))
        else:
            placeholder(s, cx, y + hh, cw, ih, sd.get("alt"))
        rect(s, cx, y + hh, cw, ih, None, "ink" if aft else "line", lw=FRAME if aft else BORDER)
        for j, r in enumerate(rows):
            ry = y + hh + ih + 0.1 + j * rh
            text(s, (cx, ry, cw * 0.6, rh), r[0], 10.5, "text", anchor="m", min_size=9,
                 name="compare label")
            text(s, (cx + cw * 0.6, ry, cw * 0.4, rh), r[1 + k], 15, "ink" if aft else "grey1",
                 bold=True, align="r", anchor="m", min_size=10, name="compare value")
            hline(s, cx, ry + rh, cw, "line", HAIR)
            if aft and len(r) > 3:
                text(s, (x + cw, ry, gut, rh), r[3], 9.5, "ink2", bold=True, align="c", anchor="m",
                     min_size=8, name="compare change")
    if p.get("delta"):
        bs = 1.15
        bx, by = x + cw + gut / 2 - bs / 2, y + hh + ih / 2 - bs / 2
        rect(s, bx, by, bs, bs, "accent" if p.get("highlight") else "ink", "paper", lw=3)
        v, _, lab = p["delta"].partition("\n")
        text(s, (bx, by + 0.14, bs, 0.5), v, 20, "paper", bold=True, align="c", anchor="m",
             min_size=12, name="delta")
        text(s, (bx + 0.05, by + 0.62, bs - 0.1, 0.36), lab, 9.5, "paper", align="c", anchor="t",
             min_size=7.5, name="delta label")


# ------------------------------------------------------------------ more panels

def p_group(s, b, p):
    """Nested grid: a panel whose body is another grid of panels."""
    grid(s, p["body"], b, p.get("gap", 0.12))


def p_icons(s, b, p):
    """Cells with a soft icon column; title and body to its right, left-aligned.
    One style: 'style' and 'align' are accepted and ignored."""
    x, y, w, h = b
    items = p["items"]
    n = len(items)
    ncol = p.get("cols", n if n <= 4 else math.ceil(n / 2))
    nrow = math.ceil(n / ncol)
    g = 0.10
    cw, chh = (w - g * (ncol - 1)) / ncol, (h - g * (nrow - 1)) / nrow
    iw = min(0.85, cw * 0.22)
    tw = cw - iw - 0.26
    ts = share([(it["title"], tw, 0.32) for it in items], 12, name="icon title",
               underfill=False, bold=True, min_size=10, line=1.05)
    bodies = [it.get("items") or it.get("text") for it in items]
    bl = any("items" in it for it in items)
    bs = share([(bd, tw, chh - 0.58) for bd in bodies], BODY_SIZE, 12.5, "icons", bullets=bl,
               underfill=not all(_lone(bd, bl) for bd in bodies if bd))
    for i, (it, body) in enumerate(zip(items, bodies)):
        cx, cy = x + (i % ncol) * (cw + g), y + (i // ncol) * (chh + g)
        hl = p.get("highlight") == i
        rect(s, cx, cy, iw, chh, "accent" if hl else "soft")
        d = min(0.40, iw * 0.6)
        icon(s, it["icon"], cx + (iw - d) / 2, cy + (chh - d) / 2, d, "paper" if hl else "ink")
        tx, ty = cx + iw + 0.14, cy + 0.12
        lone = _lone(body, "items" in it)
        if lone:  # one sentence: title + body as one centred block, not three loose pieces
            need = text_height(body, bs, tw, "items" in it)
            ty = cy + max(0.12, (chh - 0.38 - need) / 2)
        text(s, (tx, ty, tw, 0.32), it["title"], ts, "accent" if hl else "ink",
             bold=True, anchor="m", fit=False, line=1.05)
        if lone:
            text(s, (tx, ty + 0.38, tw, need + 0.05), body, bs, bullets="items" in it,
                 fit=False, name=it["title"])
        elif body:
            _body(s, (tx, cy + 0.5, tw, chh - 0.58), body, bs, "items" in it, it["title"])
        box(s, cx, cy, cw, chh, hl)


def p_venn(s, b, p):
    """2–3 outline circles; each set's title and items sit in its outer part, the centre label
    box in the overlap (highlight → accent border)."""
    x, y, w, h = b
    sets = p["sets"]
    n = len(sets)
    cols = ["ink", "ink2", "grey1"]
    cx, cy = x + w / 2, y + h / 2
    if n == 2:
        r = min(h / 2, w / 3.3)
        centers = [(cx - 0.62 * r, cy), (cx + 0.62 * r, cy)]
        mid = (cx, cy)
    else:
        r = min(h / 3.05, w / 3.25)
        cy += 0.04 * r
        centers = [(cx - 0.58 * r, cy - 0.45 * r), (cx + 0.58 * r, cy - 0.45 * r),
                   (cx, cy + 0.55 * r)]
        mid = (cx, cy - 0.12 * r)
    for (sx, sy), c in zip(centers, cols):
        rect(s, sx - r, sy - r, 2 * r, 2 * r, None, c, MSO_SHAPE.OVAL, lw=1.75)
    bw = r * (0.85 if n == 2 else 0.95)
    ts = min(15, max(12.5, 12.5 * r / 1.6))
    size = share([(st.get("items"), bw, r * 0.55) for st in sets], 10.5, max_size=12.5,
                 line=1.1, gap=2, underfill=False, name="venn items")
    for (sx, sy), c, st in zip(centers, cols, sets):
        dx, dy = sx - mid[0], sy - mid[1]
        k = math.hypot(dx, dy) or 1
        lx, ly = sx + dx / k * r * 0.42, sy + dy / k * r * 0.42
        items = st.get("items", [])
        tth = text_height(st["title"], ts, bw, bold=True, line=1.05)
        ih = text_height(items, size, bw, line=1.1, gap=2) if items else 0
        top = ly - (tth + 0.08 + ih) / 2
        text(s, (lx - bw / 2, top, bw, tth + 0.02), st["title"], ts, "ink", bold=True, align="c",
             fit=False, line=1.05)
        if items:
            text(s, (lx - bw / 2, top + tth + 0.08, bw, ih + 0.04), items, size, "text",
                 align="c", fit=False, line=1.1, gap=2)
    if p.get("center"):
        hl = bool(p.get("highlight"))
        cw = r * (0.6 if n == 2 else 0.66)
        rect(s, mid[0] - cw / 2, mid[1] - 0.36, cw, 0.72, "accent_soft" if hl else "paper",
             "accent" if hl else "ink", lw=1)
        text(s, (mid[0] - cw / 2 + 0.06, mid[1] - 0.36, cw - 0.12, 0.72), p["center"], 11, "ink",
             bold=True, align="c", anchor="m", line=1.05, gap=0, min_size=9, grow=12.5,
             name="venn center")


BMC = [("partners", "핵심 파트너", "handshake"), ("activities", "핵심 활동", "zap"),
       ("resources", "핵심 자원", "boxes"), ("value", "가치 제안", "gem"),
       ("relationships", "고객 관계", "heart-handshake"), ("channels", "채널", "megaphone"),
       ("segments", "고객 세그먼트", "users"), ("costs", "비용 구조", "receipt"),
       ("revenue", "수익원", "banknote")]


def p_canvas(s, b, p):
    """Business Model Canvas: 9 cells on one grid (column widths via 'widths'), one shared body
    size; highlighted blocks get an accent border."""
    x, y, w, h = b
    g = 0.06
    ws = p.get("widths", [1] * 5)
    cws = [(w - 4 * g) * k / sum(ws) for k in ws]
    cx = [x + sum(cws[:i]) + g * i for i in range(5)]
    th = (h - g) * p.get("split", 0.66)
    bh = h - g - th
    half = (th - g) / 2
    pos = {"partners": (cx[0], y, cws[0], th), "activities": (cx[1], y, cws[1], half),
           "resources": (cx[1], y + half + g, cws[1], half), "value": (cx[2], y, cws[2], th),
           "relationships": (cx[3], y, cws[3], half),
           "channels": (cx[3], y + half + g, cws[3], half), "segments": (cx[4], y, cws[4], th),
           "costs": (x, y + th + g, (w - g) / 2, bh),
           "revenue": (x + (w + g) / 2, y + th + g, (w - g) / 2, bh)}
    hl = p.get("highlight", [])
    hl = [hl] if isinstance(hl, str) else hl
    hh = 0.34
    jobs = [(p.get(k), pos[k][2] - 0.28, pos[k][3] - hh - 0.2) for k, _, _ in BMC]
    size = share(jobs, 10.5, max_size=11.5, bullets=True, name="canvas")
    gap = _spread_gap(jobs, size)  # one paragraph rhythm across all nine blocks
    for key, label, ic in BMC:
        bx, by, bw, bhh = pos[key]
        rect(s, bx, by, bw, hh, "head")
        box(s, bx, by, bw, bhh, key in hl)
        icon(s, ic, bx + 0.1, by + hh / 2 - 0.11, 0.22)
        text(s, (bx + 0.4, by, bw - 0.5, hh), label, 10.5, "ink", bold=True, anchor="m",
             min_size=8.5, name=label)
        if p.get(key):
            text(s, (bx + 0.14, by + hh + 0.1, bw - 0.28, bhh - hh - 0.2), p[key], size,
                 bullets=True, fit=False, gap=gap)


STATUS = {"done": ("circle-check", "ink", "완료"), "partial": ("circle-dot-dashed", "ink2", "진행"),
          "todo": ("circle", "grey2", "미착수"), "fail": ("circle-x", "ink", "미흡")}


def p_checklist(s, b, p):
    """Checklist (status icons) or scorecard ('n / 5' with a bar, optional weights), drawn like
    p_table: head-filled header, hairline rows filling the height, score row on a head strip."""
    x, y, w, h = b
    items = p["items"]
    n = len(items)
    rating = any("score" in it for it in items)
    weighted = any("weight" in it for it in items)
    cols = p.get("columns") or (["평가 항목"] + (["배점"] if weighted else []) +
                                (["점수", "근거"] if rating else ["상태", "비고"]))
    rel = p.get("widths") or ([2.6 if rating else 3.2] + ([0.8] if weighted else []) +
                              ([2.2, 3] if rating else [1.1, 3]))
    ws = [w * r / sum(rel) for r in rel]
    xs = [x + sum(ws[:i]) for i in range(len(ws))]
    hh, summary, tot = 0.36, p.get("summary"), None
    if weighted and rating and summary is None:
        tot = sum(it.get("weight", 0) * it.get("score", 0) / it.get("max", 5) for it in items)
        summary = f"/ {sum(it.get('weight', 0) for it in items)}"
    sh = 0.56 if summary else 0
    rh = (h - hh - sh) / n
    rect(s, x, y, w, hh, "head")
    for i, (cx, cw, c) in enumerate(zip(xs, ws, cols)):
        text(s, (cx + 0.12, y, cw - 0.24, hh), c, 10.5, "ink", bold=True,
             align="l" if i == 0 else "c", anchor="m", fit=False)
    if not p.get("title"):  # titled: the title rule above is our top edge
        hline(s, x, y, w, "ink", FRAME)
    hline(s, x, y + hh, w, "ink", BORDER)
    nc = len(ws) - 1
    ls = share([(it["label"], ws[0] - 0.24, rh - 0.1) for it in items], 11, 12,
               name="check label", underfill=False)
    ns = share([(it.get("note"), ws[nc] - 0.24, rh - 0.1) for it in items], 10.5, 11.5,
               name="check note", underfill=False)
    for r, it in enumerate(items):
        ry = y + hh + r * rh
        hl = p.get("highlight") == r
        if hl:
            rect(s, x, ry, w, rh, "accent_soft")
        if r:
            hline(s, x, ry, w, "line", HAIR)
        text(s, (xs[0] + 0.12, ry, ws[0] - 0.24, rh), it["label"], ls, "accent" if hl else "ink",
             bold=hl, anchor="m", fit=False)
        c = 1
        if weighted:
            text(s, (xs[1], ry, ws[1], rh), str(it.get("weight", "")), 11, "text", align="c",
                 anchor="m", fit=False)
            c = 2
        mx, mw = xs[c], ws[c]
        if rating:
            mxv = it.get("max", 5)
            bw = mw - 0.24 - 0.62
            rect(s, mx + 0.12, ry + rh / 2 - 0.05, bw, 0.10, "head")
            rect(s, mx + 0.12, ry + rh / 2 - 0.05, bw * min(1, it["score"] / mxv), 0.10,
                 "accent" if hl else "ink2")
            text(s, (mx + mw - 0.7, ry, 0.58, rh), f"**{it['score']}** / {mxv}", 10, "muted",
                 align="r", anchor="m", emph="accent" if hl else "ink", fit=False)
        else:
            st = it.get("status", "todo")
            st = "done" if st is True else "todo" if st is False else st
            ic, cl, lab = STATUS[st]
            lab = it.get("status_label", lab)
            cl = "accent" if hl else cl
            gx = mx + (mw - 0.28 - text_w(lab, 10.5, True)) / 2
            icon(s, ic, gx, ry + rh / 2 - 0.11, 0.22, cl)
            text(s, (gx + 0.28, ry, mw / 2 + 0.3, rh), lab, 10.5, cl, bold=True, anchor="m",
                 fit=False)
        if it.get("note"):
            text(s, (xs[nc] + 0.12, ry, ws[nc] - 0.24, rh), it["note"], ns,
                 "muted" if it.get("status") == "todo" else "text", anchor="m", emph="ink",
                 fit=False)
    for cx in xs[1:]:
        line(s, cx, y, cx, y + hh + n * rh, "rule", HAIR)
    if summary:
        sy = y + hh + n * rh
        rect(s, x, sy, w, sh, "head")
        hline(s, x, sy, w, "ink", BORDER)
        if tot is None:
            text(s, (x + 0.12, sy, w - 0.24, sh), summary, 11.5, "ink", bold=True, align="r",
                 anchor="m", emph="ink", name="check summary")
        else:
            text(s, (x + 0.12, sy, 2.5, sh), "종합 점수", 11, "ink", bold=True, anchor="m",
                 fit=False)
            text(s, (x + w - 0.72, sy, 0.6, sh), summary, 11, "muted", anchor="m", fit=False)
            text(s, (x + w - 2.5, sy, 1.72, sh), f"{tot:.1f}", 20, "ink", bold=True, align="r",
                 anchor="m", fit=False)
    hline(s, x, y + h, w, "ink", FRAME)


KOREA_TILES = {"서울": (1, 0), "강원": (2, 0), "인천": (0, 1), "경기": (1, 1), "충북": (2, 1),
               "경북": (3, 1), "충남": (0, 2), "세종": (1, 2), "대전": (2, 2), "대구": (3, 2),
               "전북": (0, 3), "광주": (1, 3), "경남": (2, 3), "울산": (3, 3), "전남": (1, 4),
               "부산": (2, 4), "제주": (0, 5)}


MAP_RAMP = ["head", "grey2", "grey1", "ink2"]  # choropleth, light → dark (the HEAT scale)


def _map_fill(k):
    """Colour at k in 0..1 along MAP_RAMP."""
    i = min(int(k * 3), 2)
    return mix(MAP_RAMP[i], MAP_RAMP[i + 1], k * 3 - i)


def p_map(s, b, p):
    """South Korea tile map (17 시·도): square-cornered tiles filled by value on an ink ramp,
    value in bold; highlighted 시·도 get an accent border."""
    x, y, w, h = b
    vals = p.get("values", {})
    fmt = p.get("format", "{:,}")
    hl = set(p.get("highlight", []))
    lg = 0.42
    th = (h - lg) / 6
    tw = min(w / 4, th * 1.45)
    g = 0.05
    ox, oy = x + (w - 4 * tw) / 2, y
    nums = [v for v in vals.values() if isinstance(v, (int, float))]
    lo, hi = (min(nums), max(nums)) if nums else (0, 1)
    ns, vs = min(10.5, th * 13), min(16, th * 20)
    for name, (c, r) in KOREA_TILES.items():
        v = vals.get(name)
        k = (v - lo) / ((hi - lo) or 1) if isinstance(v, (int, float)) else None
        fill = "soft" if k is None else _map_fill(k)
        tx, ty = ox + c * tw + g / 2, oy + r * th + g / 2
        if name in hl:  # accent frame inside the tile, so it never touches a neighbour
            rect(s, tx, ty, tw - g, th - g, fill)
            rect(s, tx + 0.02, ty + 0.02, tw - g - 0.04, th - g - 0.04, None, "accent", lw=2.25)
        else:
            rect(s, tx, ty, tw - g, th - g, fill, "line", lw=HAIR)
        fg = "paper" if k is not None and k > 0.55 else "ink"
        text(s, (tx + 0.08, ty + 0.05, tw - g - 0.16, th * 0.32), name, ns, fg, align="l",
             anchor="t", fit=False)
        if v is not None:
            text(s, (tx + 0.08, ty + th * 0.36, tw - g - 0.16, th * 0.55),
                 fmt.format(v) if isinstance(v, (int, float)) else str(v), vs, fg, bold=True,
                 align="r", anchor="m", fit=False)
    if nums:
        ly = y + h - lg + 0.16
        steps, sw = 5, min(0.45, w / 12)
        lx = x + (w - steps * sw) / 2
        text(s, (lx - 1.3, ly, 1.2, 0.2), fmt.format(lo), 8.5, "muted", align="r", anchor="m",
             fit=False)
        for i in range(steps):
            rect(s, lx + i * sw, ly + 0.03, sw, 0.14, _map_fill(i / (steps - 1)))
        text(s, (lx + steps * sw + 0.1, ly, 1.6, 0.2),
             f"{fmt.format(hi)} {p.get('unit') or ''}".strip(), 8.5, "muted", anchor="m",
             fit=False)


def _edge_point(n, side):
    x, y, w, h = n
    return {"l": (x, y + h / 2), "r": (x + w, y + h / 2), "t": (x + w / 2, y),
            "b": (x + w / 2, y + h)}[side]


NODE_STYLE = {"accent": ("accent", "paper", None), "ink": ("ink", "paper", None),
              "soft": ("soft", "ink", "line"), "paper": ("paper", "ink", "ink2")}


def p_diagram(s, b, p):
    """Free-form diagram on a cols×rows grid: zones, nodes and routed arrows. lanes → swimlane.
    numbered → a number cell on each node (in node order)."""
    x, y, w, h = b
    lanes = p.get("lanes")
    lw = p.get("lane_width", 1.0) if lanes else 0
    rows = len(lanes) if lanes else p.get("rows", 3)
    cols = p.get("cols", 4)
    gx, gw = x + lw, w - lw
    cw, rh = gw / cols, h / rows
    areas = []  # (box, fill) behind edge labels
    if lanes:
        for i, ln in enumerate(lanes):
            ly = y + i * rh
            fill = "soft" if i % 2 == 0 else "paper"
            rect(s, gx, ly, gw, rh, fill)
            areas.append(((gx, ly, gw, rh), fill))
            rect(s, x, ly, lw, rh, "head")
            text(s, (x + 0.08, ly, lw - 0.16, rh), ln, 10.5, "ink", bold=True, align="c",
                 anchor="m", min_size=8.5, line=1.1, name="lane")
            if i:
                hline(s, x, ly, w, "line", HAIR)
        line(s, gx, y, gx, y + h, "ink", FRAME)
        hline(s, x, y, w, "ink", FRAME)
        hline(s, x, y + h, w, "ink", FRAME)
    for z in p.get("zones", []):
        zb = _zone(s, gx + z["col"] * cw + 0.04, y + z["row"] * rh + 0.04,
                   z.get("w", 1) * cw - 0.08, z.get("h", 1) * rh - 0.08, z["label"],
                   z.get("highlight"))
        areas.append((zb, "soft"))
    pad_x, pad_top, pad_bot = 0.14, 0.34 if p.get("zones") else 0.14, 0.14
    num = p.get("numbered")
    nodes = []
    for i, nd in enumerate(p["nodes"]):
        nx = gx + nd["col"] * cw + pad_x
        ny = y + nd["row"] * rh + pad_top
        nw = nd.get("w", 1) * cw - 2 * pad_x
        nh = nd.get("h", 1) * rh - pad_top - pad_bot
        if nd.get("height"):  # fixed node height, centred in its cells
            ny += (nh - nd["height"]) / 2
            nh = nd["height"]
        lead = (0.3 if num else 0) + (min(0.34, nh * 0.5) + 0.08 if nd.get("icon") else 0)
        nodes.append((nd, (nx, ny, nw, nh), lead))
    size = share([([nd["label"]] + ([nd["sub"]] if nd.get("sub") else []),
                   (nw - lead - 0.2) * PAIR_W, nh - 0.08) for nd, (nx, ny, nw, nh), lead in nodes],
                 11, max_size=12, bold=True, min_size=9, line=1.05, gap=2, underfill=False,
                 name="diagram node")
    boxes = {}
    for i, (nd, (nx, ny, nw, nh), lead) in enumerate(nodes):
        fill, fg, border = NODE_STYLE.get(nd.get("style", "paper"), NODE_STYLE["paper"])
        rect(s, nx, ny, nw, nh, fill, border)
        tx = nx + 0.1
        if num:
            nc = rect(s, nx, ny, 0.3, nh, *(("paper", "ink") if fg == "paper" else ("ink", None)))
            shape_text(nc, i + 1, 11, "ink" if fg == "paper" else "paper")
            tx = nx + 0.38
        if nd.get("icon"):
            d = min(0.34, nh * 0.5)
            icon(s, nd["icon"], tx, ny + nh / 2 - d / 2, d, fg)
            tx += d + 0.08
        _pair(s, (tx, ny + 0.04, nx + nw - tx - 0.1, nh - 0.08), nd["label"], nd.get("sub"), size,
              fg, fg if fg == "paper" else "text", "l" if lead else "c", max(size - 1.5, 9))
        boxes[nd["id"]] = (nx, ny, nw, nh)
    for e in p.get("edges", []):
        a, z = boxes[e["from"]], boxes[e["to"]]
        c = "accent" if e.get("highlight") else "ink2"
        dash, both = e.get("style") == "dashed", e.get("both", False)
        ov_y = min(a[1] + a[3], z[1] + z[3]) - max(a[1], z[1])
        ov_x = min(a[0] + a[2], z[0] + z[2]) - max(a[0], z[0])
        if ov_y > 0.15:  # side by side → straight horizontal
            yy = max(a[1], z[1]) + ov_y / 2
            x1, x2 = (a[0] + a[2], z[0]) if a[0] < z[0] else (a[0], z[0] + z[2])
            pts = [(x1, yy), (x2, yy)]
        elif ov_x > 0.15:  # stacked → straight vertical
            xx = max(a[0], z[0]) + ov_x / 2
            y1, y2 = (a[1] + a[3], z[1]) if a[1] < z[1] else (a[1], z[1] + z[3])
            pts = [(xx, y1), (xx, y2)]
        else:  # elbow: leave sideways, enter from top/bottom
            p1 = _edge_point(a, "r" if z[0] > a[0] else "l")
            p3 = _edge_point(z, "t" if z[1] > a[1] else "b")
            pts = [p1, (p3[0], p1[1]), p3]
        for i in range(len(pts) - 1):
            line(s, *pts[i], *pts[i + 1], c, 1.5 if e.get("highlight") else 1.25,
                 arrow=i == len(pts) - 2, dash=dash, arrow_start=both and i == 0)
        if e.get("label"):
            (ax, ay), (bx2, by2) = max(zip(pts, pts[1:]),
                                       key=lambda q: math.dist(q[0], q[1]))
            mx, my = (ax + bx2) / 2, (ay + by2) / 2
            lw_ = text_w(e["label"], 9, True) + 0.16
            # ponytail: no collision check with other labels; route around if it bites
            if ay == by2 and abs(bx2 - ax) < lw_ + 0.06:  # a label off the line reads as orphaned
                warn(f"edge label {e['label']!r}: no room, widen the gap or drop it")
                continue
            if ax == bx2 and abs(by2 - ay) < 0.36:  # short vertical: beside the line
                mx += lw_ / 2 + 0.06
            my = min(max(my, y + 0.12), y + h - 0.12)
            bg = next((f for (ux, uy, uw, uh), f in reversed(areas)
                       if ux <= mx <= ux + uw and uy <= my <= uy + uh), "paper")
            lb = rect(s, mx - lw_ / 2, my - 0.12, lw_, 0.24, bg)
            shape_text(lb, e["label"], 9, c)


PANELS = {
    "bullets": p_bullets, "callout": p_callout, "table": p_table, "chart": p_chart,
    "kpi": p_kpi, "cards": p_cards, "process": p_process, "timeline": p_timeline,
    "cycle": p_cycle, "stack": p_stack, "image": p_image, "label": p_label, "arrow": p_arrow,
    "matrix": p_matrix, "pyramid": p_pyramid, "funnel": p_funnel, "tree": p_tree,
    "house": p_house, "milestones": p_milestones, "profiles": p_profiles,
    "numbered": p_numbered, "flow": p_flow, "progress": p_progress, "roadmap": p_roadmap,
    "group": p_group, "icons": p_icons, "venn": p_venn, "canvas": p_canvas,
    "checklist": p_checklist, "map": p_map, "diagram": p_diagram, "images": p_images,
    "devices": p_devices, "storyboard": p_storyboard, "compare": p_compare,
}
FIXED_W = {"arrow": 0.35, "label": 0.55}
FRAMED = {"bullets", "chart", "progress", "venn", "funnel", "pyramid", "map"}  # get a cell border
NO_OFFSET = {"arrow", "label", "group"}  # never pushed down to line up with titled siblings
# titled panels whose own top edge sits on the title rule (the rest start TITLE_H below it)
LIFT = {"cards", "kpi", "numbered", "icons", "matrix", "table", "checklist", "image", "images"}


def panel(s, b, p):
    x, y, w, h = b
    t = p.get("title")
    style = p.get("title_style", T["title_style"])
    th = 0.32 if style == "bar" else 0.30
    unit = p.get("unit") if p["type"] in ("chart", "table") else None
    note = p.get("note") or (unit and f"(단위: {unit})")
    framed = p.get("boxed") or (t and p["type"] in FRAMED)
    if framed:  # border first, so the title rule sits on top of it
        top = y + th if t else y
        rect(s, x, top, w, y + h - top, None, "line", lw=BORDER)
    if t:
        if style == "bar":
            rect(s, x, y, w, th, "ink")
            text(s, (x + 0.12, y, w - 0.24, th), t, 10.5, "paper", bold=True, anchor="m",
                 name="panel title")
        elif style == "line":
            text(s, (x, y, w, th), t, 11.5, "ink", bold=True, anchor="m", name="panel title")
            hline(s, x, y + th, w, "ink", FRAME)
        else:  # tab: ink tab sitting on a 1.5pt rule
            tw = min(w, max(1.4, text_w(t, 10.5, True) + 0.36))
            rect(s, x, y, tw, th, "ink")
            text(s, (x + 0.12, y, tw - 0.18, th), t, 10.5, "paper", bold=True, anchor="m",
                 name="panel title")
            hline(s, x, y + th, w, "ink", FRAME)
        if note:
            text(s, (x + w - 3.2, y + 0.04, 3.2, 0.22), note, 8.5, "muted", align="r", fit=False)
            if unit:
                p = {**p, "unit": None}  # drawn here, not again by the panel
        lift = th if p["type"] in LIFT and not (p.get("points") or p.get("col_labels")) \
            else TITLE_H
        y, h = y + lift, h - lift
    if framed:
        x, w = x + 0.12, w - 0.24
        y, h = (y + 0.02, h - 0.14) if t else (y + 0.12, h - 0.24)
    if p["type"] not in PANELS:
        sys.exit(f"slide {CTX['slide']}: unknown panel type {p['type']!r}. "
                 f"Choose from: {', '.join(PANELS)}")
    PANELS[p["type"]](s, (x, y, w, h), p)


def _widths(cols, w, gap):
    fixed = [FIXED_W[c["type"]] if "w" not in c and c["type"] in FIXED_W
             and not (c["type"] == "arrow" and c.get("text")) else None for c in cols]  # text strip
    flex = w - gap * (len(cols) - 1) - sum(f for f in fixed if f)
    tot = sum(c.get("w", 1) for c, f in zip(cols, fixed) if f is None) or 1
    return [f if f else flex * c.get("w", 1) / tot for c, f in zip(cols, fixed)]


def _cap(c, cw):
    """Height a panel never needs to exceed (kpi strips, callouts, arrows), or None."""
    t, extra = c["type"], TITLE_H if c.get("title") else 0
    if t == "arrow":
        return 0.34 + extra
    if t == "kpi" and c.get("direction") != "column":
        return 1.45 + extra + (0.3 if any(it.get("sub") for it in c["items"]) else 0)
    if t == "callout":
        need = text_height(c["text"], 12.5, cw - 1.45, bold=True) + 0.34
        return min(1.10, max(0.62, need)) + extra
    return None


def grid(s, body, box, gap=GAP):
    """body = [row, ...]; row = [panel, ...] or {"h": weight, "cols": [...]}; panel "w" = weight.
    Rows made only of kpi/callout/arrow panels are capped at the height they need; the surplus
    goes to the other rows. In a row with a titled panel, untitled siblings start at its title
    rule."""
    x, y, w, h = box
    rows = [r if isinstance(r, dict) else {"cols": r} for r in body]
    widths = [_widths(r["cols"], w, gap) for r in rows]
    weights = [r.get("h", 1) for r in rows]
    avail = h - gap * (len(rows) - 1)
    hs = [avail * k / sum(weights) for k in weights]
    caps = []
    for r, ws in zip(rows, widths):
        cs = [_cap(c, cw) for c, cw in zip(r["cols"], ws)]
        caps.append(max(cs) if all(cs) else None)
    if not all(caps):
        spare = sum(rh - min(rh, c) for rh, c in zip(hs, caps) if c)
        free = sum(k for k, c in zip(weights, caps) if not c)
        hs = [min(rh, c) if c else rh + spare * k / free for rh, c, k in zip(hs, caps, weights)]
    cy = y
    for r, ws, rh in zip(rows, widths, hs):
        titled = any(c.get("title") for c in r["cols"])
        cx = x
        for c, cw in zip(r["cols"], ws):
            off = 0.30 if titled and not c.get("title") and c["type"] not in NO_OFFSET else 0
            panel(s, (cx, cy + off, cw, rh - off), c)
            cx += cw + gap
        cy += rh + gap


# ------------------------------------------------------------------ slide chrome

def _section(d, deck):
    """(index, name) of the slide's section; index is None when it is not in deck['sections']."""
    sections, sec = deck.get("sections", []), d.get("section")
    if isinstance(sec, int) and sec < len(sections):
        return sec, sections[sec]
    return (sections.index(sec), sec) if sec in sections else (None, sec)


def header(s, d, deck, x=M, w=W - 2 * M, tabs=True):
    """Breadcrumb + folder tabs on a heavy rule, the headline, then the 핵심 요약 box (lead).
    x, w, tabs: column geometry (split layout). d["banner"]: image band, headline on ink."""
    sections = deck.get("sections", [])
    i, name = _section(d, deck)
    crumb = f"{ROMAN[i]}. {name}" if i is not None and i < len(ROMAN) else name
    kicker = d.get("kicker")
    if d.get("banner"):
        y = _banner(s, d["banner"], " | ".join(t for t in (crumb, kicker != name and kicker) if t),
                    d["title"])
    else:
        if sections and tabs:  # folder tabs, right-aligned on the 2pt rule
            n = len(sections)
            tw = min(max(max(text_w(nm, 9, True) for nm in sections) + 0.3, 0.95), 1.6,
                     (w - 6.0) / n - 0.02)
            tx = x + w - n * tw - (n - 1) * 0.02
            for j, nm in enumerate(sections):
                cur = j == i
                tab = rect(s, tx, 0.22, tw, 0.28, "ink" if cur else "soft",
                           None if cur else "rule", lw=HAIR)
                shape_text(tab, nm, 9, "paper" if cur else "muted", bold=cur,
                           box=(tw - 0.1, 0.28))
                tx += tw + 0.02
        if crumb:
            text(s, (x, 0.22, min(6.5, w), 0.26), crumb, 10.5, "ink", bold=True, anchor="m",
                 fit=False)
        if kicker and kicker != name:
            kx = x + (text_w(crumb, 10.5, True) + 0.12 if crumb else 0)
            text(s, (kx, 0.22, min(6.5, w), 0.26), f"|  {kicker}" if crumb else kicker, 10,
                 "muted", anchor="m", fit=False)
        hline(s, x, 0.50, w, "ink", STRONG)
        y = 0.62
        hs = fit_size(d["title"], 20, w, 0.80, bold=True, min_size=17, line=1.1, name="headline")
        hh = text_height(d["title"], hs, w, bold=True, line=1.1)
        text(s, (x, y, w, hh + 0.05), d["title"], hs, "ink", bold=True, line=1.1, fit=False)
        y += hh + 0.12
    if d.get("lead"):
        bh = max(0.46, text_height(d["lead"], 11.5, w - 1.35) + 0.16)
        rect(s, x, y, w, bh, "soft", "ink", lw=1.0)
        lab = rect(s, x, y, 1.05, bh, "ink")
        shape_text(lab, d.get("lead_label", "핵심 요약"), 10, "paper")
        text(s, (x + 1.2, y, w - 1.35, bh), d["lead"], 11.5, "text", anchor="m",
             emph="ink", min_size=10, name="lead")
        return y + bh + 0.18
    if d.get("banner"):
        return y + 0.04
    hline(s, x, y - 0.04, w, "ink", BORDER)
    return y + 0.16


def _banner(s, bn, crumb, title):
    """Content-slide banner: image band across the top, the headline on a solid ink block at its
    left. Returns the y below the band."""
    bh, bw = bn.get("h", 1.9), W * bn.get("w", 0.58)
    src = _src(bn)
    if src:
        picture(s, src, bw, 0, W - bw, bh, focus=bn.get("focus"))
    else:
        placeholder(s, bw, 0, W - bw, bh, bn.get("alt"), dark=True)
    rect(s, 0, 0, bw, bh, "ink")
    if crumb:
        text(s, (M, 0.22, bw - M - 0.35, 0.26), crumb, 10.5, "on_dark", bold=True, anchor="m",
             fit=False)
    tw = bw - M - 0.4
    hs = fit_size(title, 21, tw, bh - 0.8, bold=True, min_size=17, line=1.1, name="headline",
                  underfill=False)
    text(s, (M, 0.62, tw, bh - 0.8), title, hs, "paper", bold=True, line=1.1, fit=False,
         emph="accent_on_dark")
    if bn.get("caption"):
        cw = min(W - bw, text_w(bn["caption"], 8.5) * 1.15 + 0.3)  # margin for fallback fonts
        rect(s, W - cw, bh - 0.28, cw, 0.28, "ink")
        text(s, (W - cw + 0.15, bh - 0.28, cw - 0.2, 0.28), bn["caption"], 8.5, "on_dark",
             anchor="m", fit=False)
    return bh + 0.16


def footer(s, d, deck, n, x=M, w=W - 2 * M):
    """Source / note lines, a rule, deck footer and page number. Returns the body bottom."""
    hline(s, x, 7.12, w, "ink", BORDER)
    if deck.get("footer"):
        text(s, (x, 7.16, min(9, w - 1.2), 0.2), deck["footer"], 8, "muted", fit=False)
    text(s, (x + w - 1, 7.16, 1, 0.2), str(n), 8.5, "ink", bold=True, align="r", fit=False)
    bottom, y = 6.96, 6.90
    for key, lead in (("source", "※ 출처: "), ("note", "주: ")):
        if d.get(key):
            text(s, (x, y, w - 1.2, 0.18), lead + d[key], 8.5, "muted", fit=False)
            bottom, y = y - 0.08, y - 0.18
    return bottom


def accents(d):
    """Accent budget: headline emphasis + highlighted panels + emphasised callouts."""
    def walk(body):
        for row in body:
            for c in (row["cols"] if isinstance(row, dict) else row):
                yield c
                if c["type"] == "group":
                    yield from walk(c["body"])
    n = int("**" in d.get("title", ""))
    for c in walk(d.get("body", [])):
        hl = c.get("highlight")
        n += hl is not None and hl is not False and hl != [] or c.get("highlight_col") is not None
        n += c["type"] == "callout" and "**" in c.get("text", "")
    return n


# ------------------------------------------------------------------ full-bleed slides

def _fb_kv(items, label=""):
    """["라벨: 값", ...] → [[label, value]]; a part without ':' continues the previous value
    (or opens a row under the default label)."""
    out = []
    for it in items:
        k, sep, v = str(it).partition(":")
        if sep:
            out.append([k.strip(), v.strip()])
        elif out:
            out[-1][1] += " · " + k.strip()
        else:
            out.append([label, k.strip()])
    return out


def _fb_table(s, x, y, w, rows, rh=0.36, lw=1.4):
    """Cover-style key-value table: head label cells, 1.25pt ink top and bottom rules."""
    for i, (k, v) in enumerate(rows):
        cy = y + i * rh
        rect(s, x, cy, lw, rh, "head")
        text(s, (x, cy, lw, rh), k, 10.5, "ink", bold=True, align="c", anchor="m", fit=False)
        text(s, (x + lw + 0.15, cy, w - lw - 0.25, rh), v, 10.5, "text", anchor="m",
             min_size=9, line=1.0, name="meta")
        if i:
            hline(s, x, cy, w, "line", HAIR)
    hline(s, x, y, w, "ink", 1.25)
    hline(s, x, y + rh * len(rows), w, "ink", 1.25)


def _fb_rh(n, avail, cap, name):
    """Row height for n key-value rows in avail inches: cap, shrunk to fit, floored at 0.28."""
    rh = min(cap, avail / n)
    if rh < 0.28:
        warn(f"{name}: {n} rows do not fit above the bottom band — drop a row")
    return max(rh, 0.28)


def _fb_rule2(s, x, y, w):
    hline(s, x, y, w, "ink", 3.0)
    hline(s, x, y + 0.07, w, "ink", BORDER)


def _fb_band(s, deck):
    rect(s, 0, H - 0.42, W, 0.42, "ink")
    rect(s, 0, H - 0.42, W, 0.05, "accent")
    if deck.get("footer"):
        text(s, (0.9, H - 0.37, W - 1.8, 0.37), deck["footer"], 9, "paper", anchor="m",
             fit=False)


def _fb_pages(deck):
    """{section index: 'p. a–b'}: the first run of slides in that section, from its divider on.
    A later slide of an earlier section (an appendix, a gallery) does not stretch the range."""
    pg, last = {}, None
    for n, x in enumerate(deck["slides"], 1):
        no = str(x.get("no", "")) if x.get("layout") == "divider" else ""
        i = _section({"section": int(no) - 1} if no.isdigit() and int(no) else x, deck)[0]
        if i is None:
            continue
        if i not in pg:
            pg[i] = [n, n]
        elif pg[i][1] == last:  # still the same run
            pg[i][1] = n
        last = n
    return {i: f"p. {a}" if a == b else f"p. {a}–{b}" for i, (a, b) in pg.items()}


def _fb_toc(deck, d=None):
    """TOC rows as dicts {title, desc, items}, from the toc slide's items or deck sections.
    A plain string row takes its sub-items from the divider with the matching no."""
    d = d or next((x for x in deck["slides"] if x.get("layout") == "toc"), {})
    divs = {int(x["no"]): x for x in reversed(deck["slides"])
            if x.get("layout") == "divider" and str(x.get("no", "")).isdigit()}
    return [{**it, "title": it.get("title") or it.get("name", "")} if isinstance(it, dict)
            else {"title": it, "items": divs.get(i + 1, {}).get("items")}
            for i, it in enumerate(d.get("items") or deck.get("sections", []))]


def _fb_points(s, pts, x, y, w, h, head="근거"):
    """Evidence cells: a head strip '근거 n', then a value over its label, or a sentence."""
    n, gap = len(pts), 0.2
    pw = (w - gap * (n - 1)) / n
    pts = [p if isinstance(p, dict) else {"label": p} for p in pts]
    top = y + (0.84 if any(p.get("value") for p in pts) else 0.40)
    lh = y + h - 0.1 - top
    vs = share([(p.get("value"), pw - 0.3, 0.46) for p in pts], 22, bold=True, min_size=16,
               name="point value")
    sz = share([(p["label"], pw - 0.3, lh) for p in pts], 11, name="points")
    for i, p in enumerate(pts):
        px = x + i * (pw + gap)
        cell_head(s, px, y, pw, 0.30, f"{head} {i + 1}", size=10)
        box(s, px, y, pw, h)
        if p.get("value"):
            text(s, (px + 0.15, y + 0.38, pw - 0.3, 0.46), p["value"], vs, "ink", bold=True,
                 anchor="m", fit=False)
        text(s, (px + 0.15, top, pw - 0.3, lh), p["label"], sz, "text", fit=False, spread=True)


def f_cover(s, d, deck):
    rect(s, 0, 0, W, H, "paper")
    src = _src({"src": d.get("image")})
    tw, rw = (7.2, W - 5.1 - 0.9) if src else (11.5, W - 1.8)
    if src:  # photo on the right, above the band
        picture(s, src, W - 4.6, 0, 4.6, H - 0.42, d.get("crop", "center"), d.get("focus"))
    ts = fit_size(d["title"], 34, tw, 2.1, bold=True, line=1.12, min_size=26,
                  max_size=40 if src else 46, underfill=False, name="cover title")
    th = text_height(d["title"], ts, tw, bold=True, line=1.12)
    if d.get("kicker"):  # sits on the title, white space goes above it
        kw = text_w(d["kicker"], 11, True) + 0.4
        shape_text(rect(s, 0.9, 3.9 - th - 0.62, kw, 0.36, None, "ink"), d["kicker"], 11, "ink")
    text(s, (0.9, 1.8, tw, 2.1), d["title"], ts, "ink", bold=True, anchor="b", line=1.12,
         fit=False)
    _fb_rule2(s, 0.9, 4.10, rw)
    if d.get("subtitle"):
        text(s, (0.9, 4.32, tw, 0.75), d["subtitle"], 14, "muted", min_size=11,
             name="cover subtitle")
    meta = d.get("meta")
    rows = _fb_kv(meta) if meta else (
        (_fb_kv(d["org"].split(" · "), "기관") if d.get("org") else [])
        + (_fb_kv(d["author"].split(" · "), "연구책임자") if d.get("author") else [])
        + ([["제출일", d["date"]]] if d.get("date") else []))
    orgs = _fb_kv(d["org"].split(" · "), "기관") if meta and d.get("org") else []
    if src:  # no room for the org lockup beside the table
        rows, orgs = rows + orgs, []
    y0, avail = 5.2, H - 0.42 - 0.12 - 5.2
    if rows and not orgs:
        mw = rw if src else 5.6
        _fb_table(s, 0.9 if src else W - 0.9 - mw, y0, mw, rows,
                  _fb_rh(len(rows), avail, 0.36, "cover meta"))
    elif orgs:  # 주관·공동기관 beside the table, both blocks the same height
        tot = min(max(0.36 * len(rows), 0.62 * len(orgs)), avail)
        if rows:
            _fb_table(s, W - 0.9 - 5.6, y0, 5.6, rows, _fb_rh(len(rows), tot, 1, "cover meta"))
        ow, oh = W - 1.8 - (5.6 + 0.6 if rows else 0), tot / len(orgs)
        for i, (k, v) in enumerate(orgs):
            oy = y0 + i * oh
            if i:
                hline(s, 0.9, oy, ow, "line", HAIR)
            text(s, (0.9, oy + 0.07, ow, 0.24), k, 10, "ink", bold=True, fit=False)
            text(s, (0.9, oy + 0.31, ow, oh - 0.35), v, 12.5, "ink", bold=True, anchor="m",
                 min_size=10, line=1.05, name="cover org")
        hline(s, 0.9, y0, ow, "ink", 1.25)
        hline(s, 0.9, y0 + tot, ow, "ink", 1.25)
    _fb_band(s, deck)


def _fb_figure(s, d, deck, cur, src, x, y, w, h):
    """Right-side image of a divider: bleeds when its shape is close to the box, else sits
    whole on a soft figure cell with a border, a [그림] caption and, in the height the image
    leaves, a one-row section strip (fit: cover | contain)."""
    ia = aspect(src)
    fit = d.get("fit") or ("contain" if abs(ia / (w / h) - 1) > 0.25 else "cover")
    if fit != "contain":
        picture(s, src, x, y, w, h, d.get("crop", "center"), d.get("focus"))
        return
    rect(s, x, y, w, h, "soft")
    secs, pages = deck.get("sections", []), _fb_pages(deck)
    cap = CAP_H if d.get("caption") else 0
    iw = w - 2 * M
    ih = min(iw / ia, h - 2 * M - cap)
    sh = 0.84 if secs and len(secs) <= 6 and h - 2 * M - cap - ih >= 1.2 else 0
    iw = ih * ia  # ponytail: strip only for ≤6 sections under a wide image
    ix, iy = x + (w - iw) / 2, y + (h - ih - cap - (sh + 0.36 if sh else 0)) / 2
    picture(s, src, ix, iy, iw, ih)
    rect(s, ix, iy, iw, ih, None, "line", lw=BORDER)
    if cap:
        caption(s, ix, iy + ih, iw, d["caption"])
    if not sh:
        return
    sy, n = iy + ih + cap + 0.36, len(secs)
    cw = (iw - 0.12 * (n - 1)) / n
    for j, nm in enumerate(secs):
        cx, me = ix + j * (cw + 0.12), j == cur
        box(s, cx, sy, cw, sh, fill="paper")
        shape_text(rect(s, cx, sy, 0.44, sh, "ink" if me else "head"),
                   ROMAN[j] if j < len(ROMAN) else str(j + 1), 12, "paper" if me else "muted")
        text(s, (cx + 0.56, sy + 0.1, cw - 0.64, 0.36), nm, 12, "ink" if me else "muted",
             bold=me, anchor="m", min_size=9, name="divider strip")
        if j in pages:
            text(s, (cx + 0.56, sy + 0.46, cw - 0.64, 0.26), pages[j], 9.5,
                 "ink" if me else "muted", bold=me, anchor="m", fit=False)


def f_divider(s, d, deck):
    rect(s, 0, 0, W, H, "paper")
    rect(s, 0, 0, 5.0, H, "ink")
    no = str(d.get("no", ""))
    cur = int(no) - 1 if no.isdigit() else None
    num = no if cur is None else ROMAN[cur] if 0 <= cur < len(ROMAN) else ""
    if deck.get("title"):
        text(s, (0.9, 0.6, 3.8, 0.3), deck["title"], 10, "on_dark", bold=True, fit=False)
    if num:  # a number cell, the same device as the section map's
        shape_text(rect(s, 0.9, 1.6, 1.1, 1.1, None, "paper", lw=FRAME), num, 54, "paper",
                   box=(1.0, 1.0))
    rect(s, 0.9, 2.95, 0.7, 0.06, "accent")
    ts = fit_size(d["title"], 28, 3.8, 1.4, bold=True, line=1.1, min_size=20,
                  name="divider title")
    th = text_height(d["title"], ts, 3.8, bold=True, line=1.1)
    text(s, (0.9, 3.15, 3.8, th + 0.05), d["title"], ts, "paper", bold=True, line=1.1,
         fit=False)
    if d.get("items"):
        hy = 3.15 + th + 0.3
        hline(s, 0.9, hy, 3.8, "on_dark_rule", HAIR)
        ih = 6.7 - hy - 0.2
        sz = fit_size(d["items"], 12.5, 3.8, ih, True, gap=6, min_size=10, max_size=14,
                      underfill=False, name="divider items")  # a short agenda is fine here
        text(s, (0.9, hy + 0.2, 3.8, ih), d["items"], sz, "on_dark", bullets=True, gap=6,
             fit=False, spread=True)
    src = _src({"src": d.get("image")})
    if src:  # the image takes the right side, no veil, no section map
        _fb_figure(s, d, deck, cur, src, 5.0, 0, W - 5.0, H)
        return
    secs, pages = deck.get("sections", []), _fb_pages(deck)
    if not secs:
        return
    x0, x1 = 5.7, W - M
    hline(s, x0, 7.12, x1 - x0, "ink", BORDER)  # footer, white side only
    text(s, (x1 - 1, 7.16, 1, 0.2), str(CTX["slide"]), 8.5, "ink", bold=True, align="r",
         fit=False)
    descs = [it.get("desc") for it in _fb_toc(deck)]
    n, y0 = len(secs), 1.14
    rh = min(1.3 if any(descs) else 0.95, 5.8 / n)
    text(s, (x0, 0.6, 4, 0.3), "목 차", 13, "ink", bold=True, fit=False)  # level with deck title
    hline(s, x0, 1.0, x1 - x0, "ink", STRONG)
    cs = min(0.56, rh - 0.14)
    for j, nm in enumerate(secs):
        cy, me = y0 + j * rh, j == cur
        desc = descs[j] if j < len(descs) and rh >= 0.9 else None
        c = rect(s, x0, cy + (rh - cs) / 2, cs, cs, "ink" if me else "soft",
                 None if me else "rule")
        shape_text(c, ROMAN[j] if j < len(ROMAN) else str(j + 1), 14, "paper" if me else "muted",
                   box=(cs, cs))
        tx, ty = x0 + cs + 0.24, cy + (rh - (0.72 if desc else 0.4)) / 2
        text(s, (tx, ty, x1 - tx - 1.3, 0.4), nm, 17, "ink" if me else "muted", bold=me,
             anchor="m", fit=False)
        if desc:
            text(s, (tx, ty + 0.42, x1 - tx - 1.3, 0.3), desc, 11, "text" if me else "muted",
                 anchor="m", min_size=9, name="divider map")
        if j in pages:
            text(s, (x1 - 1.2, ty, 1.2, 0.4), pages[j], 11, "ink" if me else "muted",
                 bold=me, align="r", anchor="m", fit=False)
        hline(s, x0, cy + rh, x1 - x0, "ink" if me else "rule", STRONG if me else HAIR)


def f_toc(s, d, deck):
    rect(s, 0, 0, W, H, "paper")
    items, pages = _fb_toc(deck, d), _fb_pages(deck)
    text(s, (M, 0.55, 6, 0.5), d.get("title", "목 차"), 24, "ink", bold=True, anchor="m",
         fit=False)
    hline(s, M, 1.12, W - 2 * M, "ink", STRONG)
    hline(s, M, 1.17, W - 2 * M, "ink", HAIR)
    ncol = 1 if len(items) <= 6 else 2
    per = math.ceil(len(items) / ncol)
    cw = (W - 2 * M - 0.4 * (ncol - 1)) / ncol
    rh = min(1.35 if any(it.get("desc") or it.get("items") for it in items) else 0.95, 5.4 / per)
    for i, it in enumerate(items):
        cx, cy = M + (i // per) * (cw + 0.4), 1.45 + (i % per) * rh
        shape_text(rect(s, cx, cy + 0.12, 0.6, rh - 0.24, "ink"),
                   ROMAN[i] if i < len(ROMAN) else str(i + 1), 16, "paper")
        tx, pg = cx + 0.8, pages.get(i, "")
        text(s, (tx, cy + 0.12, cw - 2.0, 0.44), it["title"], 18, "ink", bold=True, anchor="m",
             min_size=13, name="toc item")
        pw = text_w(pg, 12, True) if pg else -0.1
        lx = tx + min(text_w(it["title"], 18, True), cw - 2.0) + 0.2
        if lx < cx + cw - pw - 0.25:  # the leader runs even without a page, for the rhythm
            line(s, lx, cy + 0.36, cx + cw - pw - 0.2, cy + 0.36, "rule", 1.0, dash=True)
        if pg:
            text(s, (cx + cw - pw - 0.1, cy + 0.12, pw + 0.1, 0.44), pg, 12, "ink2", bold=True,
                 align="r", anchor="m", fit=False)
        sy = cy + 0.62
        if it.get("desc") and rh > 0.8:
            text(s, (tx, sy, cw - 2.0, 0.26), it["desc"], 11, "text", anchor="m", min_size=9,
                 name="toc desc")
            sy += 0.3
        if it.get("items") and cy + rh - 0.1 - sy > 0.2:
            text(s, (tx, sy, cw - 2.0, cy + rh - 0.1 - sy), " · ".join(it["items"]), 10.5,
                 "muted", min_size=9, name="toc sub-items")
        hline(s, cx, cy + rh, cw, "line", HAIR)
    footer(s, d, deck, CTX["slide"])


def f_closing(s, d, deck):
    """감사합니다 + double rule + message, optional goal cells (points), contact table."""
    rect(s, 0, 0, W, H, "paper")
    pts = d.get("points", [])[:4]
    y, cw = (1.05, 8.0) if pts else (2.2, 6.0)  # points need the room: lift the stack
    text(s, (0, y, W, 0.9), d.get("title", "감사합니다"), 32, "ink", bold=True, align="c",
         anchor="m", min_size=24, name="closing title")
    _fb_rule2(s, W / 2 - cw / 2, y + 1.1, cw)
    if d.get("subtitle"):
        text(s, (1.0, y + 1.3, W - 2.0, 0.5), d["subtitle"], 13, "text", align="c", anchor="m",
             min_size=11, name="closing sub")
    y += 2.05
    if pts:
        _fb_points(s, pts, W / 2 - cw / 2, y, cw, 1.4, "목표")
        y += 1.75
    if d.get("items"):
        rows = _fb_kv(d["items"])
        _fb_table(s, W / 2 - cw / 2, y, cw, rows,
                  _fb_rh(len(rows), H - 0.42 - 0.15 - y, 0.38, "closing items"))
    _fb_band(s, deck)


def f_statement(s, d, deck):
    rect(s, 0, 0, W, H, "paper")
    i, name = _section(d, deck)
    crumb = f"{ROMAN[i]}. {name}" if i is not None else deck.get("title", "")
    text(s, (M, 0.22, 8, 0.26), crumb, 10.5, "ink", bold=True, anchor="m", fit=False)
    hline(s, M, 0.50, W - 2 * M, "ink", STRONG)
    rect(s, 1.0, 1.45, W - 2.0, 3.6, None, "ink", lw=2.25)
    rect(s, 1.07, 1.52, W - 2.14, 3.46, None, "ink", lw=HAIR)
    k = d.get("kicker", "핵심 메시지")
    shape_text(rect(s, 1.4, 1.25, text_w(k, 11.5, True) + 0.5, 0.40, "ink"), k, 11.5, "paper")
    tw, sub = W - 3.2, d.get("subtitle")
    ts = fit_size(d["title"], 26, tw, 2.2, bold=True, line=1.15, min_size=20, max_size=34,
                  underfill=False, name="statement")
    th = text_height(d["title"], ts, tw, bold=True, line=1.15)
    ss = fit_size(sub, 13, tw, 0.8, min_size=11, name="statement sub") if sub else 0
    sh = text_height(sub, ss, tw) if sub else 0
    y = 1.52 + (3.46 - th - (0.37 + sh if sub else 0)) / 2  # claim + rule + subtitle, centred
    text(s, (1.6, y, tw, th), d["title"], ts, "ink", bold=True, align="c", anchor="m",
         line=1.15, fit=False)
    if sub:
        hline(s, W / 2 - 0.6, y + th + 0.25, 1.2, "ink", BORDER)
        text(s, (1.6, y + th + 0.37, tw, sh), sub, ss, "text", align="c", anchor="m", fit=False)
    if d.get("points"):
        _fb_points(s, d["points"][:4], 1.0, 5.30, W - 2.0, 1.5)
    footer(s, d, deck, CTX["slide"])


def _stats(s, x, y, w, h, stats, dark=True, vsize=22):
    """Key figures as ruled rows (value left, label right) for photo / split columns."""
    n = len(stats)
    rh = h / n
    vs = vsize
    while vs > 14 and max(text_w(st["value"], vs, True) for st in stats) > w * 0.5:
        vs -= 1
    vw = max(text_w(st["value"], vs, True) for st in stats) + 0.25
    rule, fg = ("on_dark_rule", "paper") if dark else ("line", "ink")
    for i, st in enumerate(stats):
        ry = y + i * rh
        hline(s, x, ry, w, rule, HAIR)
        text(s, (x, ry, vw, rh), st["value"], vs, fg, bold=True, anchor="m", fit=False)
        text(s, (x + vw, ry, w - vw, rh), st.get("label", ""), 10.5, "on_dark" if dark else "text",
             anchor="m", min_size=8.5, line=1.1, name="stat label")
    hline(s, x, y + h, w, rule, HAIR)


def _tag(s, x, y, t, fg="paper", border="on_dark"):
    """Bordered kicker tag (cover style); returns its height."""
    rect(s, x, y, text_w(t, 10.5, True) + 0.3, 0.34, None, border, lw=BORDER)
    text(s, (x + 0.15, y, text_w(t, 10.5, True) + 0.1, 0.34), t, 10.5, fg, bold=True, anchor="m",
         fit=False)
    return 0.34


def f_photo(s, d, deck):
    """Photo or screen beside a solid ink text column (style 'split', default) or above an ink
    band ('band'). Kicker, title, subtitle and points stack from their real heights."""
    src = _src({"src": d.get("image")})
    pts = d.get("points", [])[:4]
    band = d.get("style") == "band"
    right = d.get("align") == "right"
    ib = (0, 0, W, 5.1) if band else (0 if right else 4.9, 0, W - 4.9, H)
    if src:
        picture(s, src, *ib, d.get("crop", "center"), focus=d.get("focus"))
    else:
        placeholder(s, *ib, d.get("alt"), dark=True)
    sub = d.get("subtitle")

    def stack(tx, tw, y0, y1, size, th_max, sub, pts):  # kicker, title, subtitle, points: centred
        th_max = min(th_max, 2 * size * 1.1 * 1.2 / 72 + 0.05)  # at most 2 lines
        ts = fit_size(d["title"], size, tw, th_max, bold=True, min_size=size - 6, line=1.1,
                      name="photo title", underfill=False)
        th = text_height(d["title"], ts, tw, bold=True, line=1.1)
        sh = text_height(sub, 12.5, tw, line=1.2) if sub else 0
        blk = (0.48 if d.get("kicker") else 0) + th + (0.24 + sh if sub else 0) + \
            (0.4 + 0.62 * len(pts) if pts else 0)
        y = y0 + max(0, (y1 - y0 - blk) / 2)
        if d.get("kicker"):
            y += _tag(s, tx, y, d["kicker"]) + 0.14
        text(s, (tx, y, tw, th + 0.05), d["title"], ts, "paper", bold=True, line=1.1, fit=False,
             emph="accent_on_dark")
        y += th + 0.24
        if sub:
            text(s, (tx, y, tw, sh + 0.05), sub, 12.5, "on_dark", line=1.2, fit=False)
            y += sh + 0.4
        if pts:
            _stats(s, tx, y, tw, 0.62 * len(pts), pts, vsize=20)

    if band:
        rect(s, 0, 5.1, W, H - 5.1, "ink")
        stack(M, 8.1 if pts or sub else W - 2 * M, 5.1, H, 26, 1.3, sub if pts else None, [])
        if pts:
            _stats(s, 8.9, 5.35, W - M - 8.9, H - 0.3 - 5.35, pts, vsize=20)
        elif sub:
            text(s, (8.9, 5.35, W - M - 8.9, H - 0.3 - 5.35), sub, 12, "on_dark", anchor="m",
                 min_size=10, name="photo sub")
    else:
        px = W - 4.9 if right else 0
        rect(s, px, 0, 4.9, H, "ink")
        line(s, 4.9 if not right else px, 0, 4.9 if not right else px, H, "on_dark_rule", BORDER)
        stack(px + 0.7, 3.6, 0.6, H - 0.6, 24, 2.4, sub, pts)
    if d.get("caption"):  # solid ink tab on the image edge that meets the ink column / band
        cw = text_w(d["caption"], 9) * 1.15 + 0.3
        cx = W - M - cw if band else (ib[0] if not right else ib[0] + ib[2] - cw)
        cy = 5.1 - 0.32 if band else H - 0.32
        rect(s, cx, cy, cw, 0.32, "ink")
        text(s, (cx + 0.15, cy, cw - 0.2, 0.32), d["caption"], 9, "paper", anchor="m", fit=False)


def f_split(s, d, deck):
    """Half-bleed image (or a framed screen on an ink / soft mat, with key figures) on one side,
    a normal header + body grid on the other."""
    rect(s, 0, 0, W, H, "paper")
    iw = W * min(0.5, max(0.36, d.get("ratio", 0.42)))
    left = d.get("side", "left") == "left"
    ix = 0 if left else W - iw
    src, mat, frame = _src({"src": d.get("image")}), d.get("mat"), d.get("frame")
    stats, cap, focus = d.get("stats", []), d.get("caption"), d.get("focus")
    if mat:
        dark = mat == "ink"
        rect(s, ix, 0, iw, H, mat)
        room = H - 1.1 - (0.34 if cap else 0) - (0.5 + 0.6 * len(stats) if stats else 0)
        a = {"laptop": 1.76, "phone": 9 / 19.5}.get(frame) or (aspect(src, focus) if src else 1.6)
        extra = {"laptop": 0.09, "browser": 0.26}.get(frame, 0)  # base / title strip
        sw = min(iw - 1.0, (room - extra) * a)  # too tall → narrower, never a wider crop
        sh = sw / a + extra
        sy = 0.6 if stats else (H - sh - (0.34 if cap else 0)) / 2
        o = _frame(s, (ix + (iw - sw) / 2, sy, sw, sh), frame, src, d.get("alt"), focus,
                   d.get("screen"), cell=False)[2]
        sx, sw, y = o[0], o[2], o[1] + o[3]
        if cap:
            text(s, (sx, y + 0.1, sw, 0.24), cap, 9, "on_dark" if dark else "muted", fit=False)
            y += 0.34
        if stats:
            _stats(s, sx, y + 0.4, sw, H - 0.55 - y - 0.4, stats, dark, vsize=24)
    else:
        if src:
            picture(s, src, ix, 0, iw, H, focus=focus)
        else:
            placeholder(s, ix, 0, iw, H, d.get("alt"), dark=True)
        sh = 0.3 + 0.62 * len(stats) if stats else 0
        if stats:
            rect(s, ix, H - sh, iw, sh, "ink")
            _stats(s, ix + 0.4, H - sh + 0.15, iw - 0.8, sh - 0.3, stats, vsize=22)
        if cap:
            cw = text_w(cap, 9) * 1.15 + 0.3
            cx = ix + iw - cw if left else ix
            rect(s, cx, H - sh - 0.32, cw, 0.32, "ink")
            text(s, (cx + 0.15, H - sh - 0.32, cw - 0.2, 0.32), cap, 9, "paper", anchor="m",
                 fit=False)
    cx = iw + 0.45 if left else M
    cw = (W - M - cx) if left else (W - iw - 0.45 - M)
    top = header(s, d, deck, cx, cw, tabs=False)
    bottom = footer(s, d, deck, CTX["slide"], cx, cw)
    grid(s, d.get("body", []), (cx, top, cw, bottom - top))
    if accents(d) >= 4:
        warn(f"{accents(d)} accent highlights — keep the headline emphasis plus one highlight "
             "(target ≤2)")


FULL = {"cover": f_cover, "divider": f_divider, "toc": f_toc, "statement": f_statement,
        "photo": f_photo, "closing": f_closing, "split": f_split}


# ------------------------------------------------------------------ build

def build(deck, out, base=Path(".")):
    T.clear()
    T.update(THEME, **deck.get("theme", {}))
    derive()
    WARN.clear()
    CTX.update(base=base, fig=0)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for n, d in enumerate(deck["slides"], 1):
        CTX["slide"] = n
        s = prs.slides.add_slide(prs.slide_layouts[6])
        lay = d.get("layout", "content")
        if lay in FULL:
            FULL[lay](s, d, deck)
        elif lay == "content":
            rect(s, 0, 0, W, H, "paper")
            top = header(s, d, deck)
            bottom = footer(s, d, deck, n)
            grid(s, d.get("body", []), (M, top, W - 2 * M, bottom - top))
            if accents(d) >= 4:
                warn(f"{accents(d)} accent highlights — keep the headline emphasis plus one "
                     "highlight (target ≤2)")
        else:
            sys.exit(f"slide {n}: unknown layout {lay!r}. Use content or one of {', '.join(FULL)}")
        if d.get("notes"):
            s.notes_slide.notes_text_frame.text = d["notes"]
    prs.save(out)
    return len(prs.slides)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("-o", "--out", default="deck.pptx")
    a = ap.parse_args()
    path = Path(a.deck)
    n = build(json.loads(path.read_text(encoding="utf-8")), a.out, path.parent)
    for w_ in WARN:
        print("WARN", w_, file=sys.stderr)
    print(f"wrote {a.out} ({n} slides, {len(WARN)} warnings)")
